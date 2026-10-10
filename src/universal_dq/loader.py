"""Read a CSV or one Excel worksheet from memory.

The bytes are parsed and then discarded by the caller. This module does not
write the upload to disk and does not log cell values.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from zipfile import BadZipFile

import pandas as pd
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.reader.excel import load_workbook

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TOTAL_UPLOAD_BYTES = 60 * 1024 * 1024
_SUPPORTED = ("csv", "xlsx")


class FileLoadError(Exception):
    """The upload could not be read as a CSV or Excel worksheet."""


@dataclass(frozen=True)
class UploadInspection:
    """File kind and, for Excel, the worksheet names."""

    file_format: str
    sheet_names: tuple[str, ...]


@dataclass(frozen=True)
class LoadedTable:
    """One table held only in memory."""

    frame: pd.DataFrame
    source_name: str
    file_format: str
    sheet_name: str | None
    sheet_names: tuple[str, ...]


def inspect_upload(
    content: bytes,
    filename: str,
    *,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> UploadInspection:
    """Check the size and format. For Excel, return worksheet names only."""
    file_format = _validated_format(content, filename, max_bytes)
    if file_format == "csv":
        _decode_csv(content)
        return UploadInspection("csv", ())
    names = _excel_sheet_names(content)
    return UploadInspection("xlsx", names)


def load_tabular_file(
    content: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> LoadedTable:
    """Load one CSV, or one selected Excel worksheet, into a DataFrame."""
    file_format = _validated_format(content, filename, max_bytes)
    source_name = _safe_name(filename)
    if file_format == "csv":
        if sheet_name is not None:
            raise FileLoadError("Worksheets apply to Excel files. This upload is a CSV.")
        columns, rows = _read_csv(content)
        frame = pd.DataFrame(rows, columns=columns, dtype=object)
        return LoadedTable(frame, source_name, "csv", None, ())

    names = _excel_sheet_names(content)
    selected = names[0] if sheet_name is None else sheet_name
    if selected not in names:
        available = ", ".join(names)
        raise FileLoadError(
            f"Worksheet {selected!r} was not found. Available worksheets: {available}."
        )
    columns, rows = _read_excel_sheet(content, selected)
    frame = pd.DataFrame(rows, columns=columns, dtype=object)
    return LoadedTable(frame, source_name, "xlsx", selected, names)


def _validated_format(content: bytes, filename: str, max_bytes: int) -> str:
    if max_bytes < 1:
        raise FileLoadError("The upload limit must be at least 1 byte.")
    if len(content) == 0:
        raise FileLoadError("The file is empty. Upload a CSV or Excel file with a header row.")
    if len(content) > max_bytes:
        limit_mb = max_bytes / (1024 * 1024)
        raise FileLoadError(
            f"The file is larger than the {limit_mb:g} MB upload limit. "
            "Choose a smaller CSV or Excel file."
        )
    suffix = _suffix(filename)
    if suffix not in _SUPPORTED:
        raise FileLoadError(
            "This file format is not supported. Upload a .csv or .xlsx file."
        )
    return suffix


def _suffix(filename: str) -> str:
    name = _safe_name(filename).lower()
    if "." not in name:
        return ""
    return name.rsplit(".", 1)[1]


def _safe_name(filename: str) -> str:
    cleaned = filename.replace("\\", "/").split("/")[-1].strip()
    return cleaned or "upload"


def _decode_csv(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileLoadError(
            "This CSV is not valid UTF-8. Save it as UTF-8 and upload it again."
        ) from exc


def _read_csv(content: bytes) -> tuple[list[str], list[list[object]]]:
    text = _decode_csv(content).replace("\r\n", "\n").replace("\r", "\n")
    rows = list(csv.reader(StringIO(text)))
    if not rows or all(cell.strip() == "" for cell in rows[0]):
        raise FileLoadError("The CSV has no header row. Put column names in the first row.")
    headers = _validate_headers(rows[0])
    width = len(headers)
    data: list[list[object]] = []
    for offset, raw in enumerate(rows[1:], start=2):
        if len(raw) > width:
            raise FileLoadError(
                f"Row {offset} has more values than the header has columns."
            )
        padded = list(raw) + [""] * (width - len(raw))
        data.append(padded)
    return headers, data


def _excel_sheet_names(content: bytes) -> tuple[str, ...]:
    workbook = _open_workbook(content)
    try:
        names = tuple(workbook.sheetnames)
    finally:
        workbook.close()
    if not names:
        raise FileLoadError("The Excel file has no worksheets.")
    return names


def _read_excel_sheet(content: bytes, sheet_name: str) -> tuple[list[str], list[list[object]]]:
    workbook = _open_workbook(content)
    try:
        worksheet = workbook[sheet_name]
        grid = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not grid:
        raise FileLoadError(
            f"Worksheet {sheet_name!r} has no header row. Put column names in the first row."
        )
    headers = _validate_headers([_header_text(cell) for cell in grid[0]])
    width = len(headers)
    data: list[list[object]] = []
    for offset, raw in enumerate(grid[1:], start=2):
        values = list(raw)
        if len(values) > width:
            raise FileLoadError(
                f"Row {offset} has more cells than the header has columns."
            )
        values.extend([None] * (width - len(values)))
        data.append(values[:width])
    return headers, data


def _open_workbook(content: bytes):
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise FileLoadError(
            "This Excel file could not be opened. It may be damaged or password protected."
        ) from exc


def _header_text(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, str):
        return cell
    return str(cell)


def _validate_headers(raw_headers: list[str]) -> list[str]:
    if not raw_headers:
        raise FileLoadError("The header row is empty. Put a name in every column.")
    headers = [str(name).strip() for name in raw_headers]
    blanks = [str(index) for index, name in enumerate(headers, start=1) if name == ""]
    if blanks:
        positions = ", ".join(blanks)
        raise FileLoadError(
            "The header row has a blank name in column position "
            f"{positions}. Add a name for every column and upload the file again."
        )
    seen: set[str] = set()
    duplicates: list[str] = []
    for name in headers:
        if name in seen and name not in duplicates:
            duplicates.append(name)
        seen.add(name)
    if duplicates:
        names = ", ".join(duplicates)
        raise FileLoadError(
            f"The header row repeats these column names: {names}. "
            "Use each column name once and upload the file again."
        )
    return headers
