"""In-memory CSV and Excel downloads for cleaned copies and the audit.

Numeric, date, and boolean cells stay typed in Excel. Formula-like text is
escaped. The audit workbook records counts and settings, not fill constants
or cell values.
"""

from __future__ import annotations

import csv
import math
from datetime import date, datetime
from io import BytesIO, StringIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from universal_dq.cleaning import CleaningSession
from universal_dq.models import PROCESSING_NOTE, AuditEntry, BatchReport, RelationshipSnapshot
from universal_dq.profile import is_missing
from universal_dq.report import escape_excel_text

CLEANING_LIMITATIONS = (
    "Approved cleaning changes a working copy. The original upload stays in memory and source files are not overwritten.",
    "Uploads are not saved into the project and are not sent to an external model.",
    "Exact duplicate removal keeps the first matching row and uses the same comparison as the quality checks.",
    "The literal text NA and NULL is kept. Empty cells and whitespace-only strings are missing.",
    "Median fill runs only when every non-missing value in the column is already numeric. The median is not written in this audit.",
    "Dates are parsed only with the format you enter. Ambiguous dates are not guessed. A value that does not match is left unchanged.",
    "Failed type conversions are counted and the original cell is kept.",
    "Candidate keys and relationships are suggestions. Orphan references are not deleted or repaired.",
    "This audit records counts, column names, and operation settings. It does not record fill constants or cell values.",
    "Cleaned rows are in the separate cleaned-dataset download.",
)

AUDIT_SHEETS = (
    "Summary",
    "Cleaned Dataset Inventory",
    "Before After Metrics",
    "Relationship Validation",
    "Cleaning Audit",
    "Limitations",
)


def cleaned_dataset_csv(frame: pd.DataFrame) -> bytes:
    """Serialize one working copy. Text that could start a formula is escaped."""
    buffer = StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([_csv_text(name) for name in frame.columns])
    for row in frame.itertuples(index=False, name=None):
        writer.writerow([_csv_text(value) for value in row])
    return buffer.getvalue().encode("utf-8-sig")


def cleaned_datasets_xlsx(session: CleaningSession) -> bytes:
    """Write every working copy into one workbook, in memory."""
    workbook = Workbook()
    inventory = workbook.active
    inventory.title = "Dataset Inventory"
    _write_rows(
        inventory,
        ["Sheet", "Dataset", "File", "Worksheet", "Rows", "Columns"],
        [
            [
                f"Data {index}",
                item.dataset_id,
                item.source_name,
                item.sheet_name or "",
                int(len(item.working)),
                int(len(item.working.columns)),
            ]
            for index, dataset_id in enumerate(session.order, start=1)
            for item in [session.copies[dataset_id]]
        ],
    )
    for index, dataset_id in enumerate(session.order, start=1):
        item = session.copies[dataset_id]
        worksheet = workbook.create_sheet(f"Data {index}")
        _write_frame(worksheet, item.working)
    return _save(workbook)


def cleaning_audit_xlsx(
    session: CleaningSession,
    original: BatchReport,
    current: BatchReport,
) -> bytes:
    """Write the cleaning audit. Frames and fill constants are not included."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, session, original, current)
    _write_inventory(workbook.create_sheet("Cleaned Dataset Inventory"), session)
    _write_metrics(workbook.create_sheet("Before After Metrics"), original, current)
    _write_relationships(workbook.create_sheet("Relationship Validation"), original, current)
    _write_audit(workbook.create_sheet("Cleaning Audit"), session.audit)
    _write_rows(
        workbook.create_sheet("Limitations"),
        ["Limitation"],
        [[note] for note in CLEANING_LIMITATIONS],
    )
    return _save(workbook)


def cell_for_excel(value: object) -> object:
    """Preserve typed values and escape untrusted text."""
    if is_missing(value):
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        number = float(value)
        if not math.isfinite(number):
            return None
        return number
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value
    text = value if isinstance(value, str) else str(value)
    return escape_excel_text(text)


def _csv_text(value: object) -> str:
    exported = cell_for_excel(value)
    if exported is None:
        return ""
    if isinstance(exported, bool):
        return "True" if exported else "False"
    if isinstance(exported, int):
        return str(exported)
    if isinstance(exported, float):
        return str(exported)
    if isinstance(exported, datetime):
        return exported.isoformat(sep=" ")
    if isinstance(exported, date):
        return exported.isoformat()
    return str(exported)


def _write_frame(worksheet: Worksheet, frame: pd.DataFrame) -> None:
    headers = [escape_excel_text(str(name)) for name in frame.columns]
    for column, header in enumerate(headers, start=1):
        worksheet.cell(1, column, header)
    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            worksheet.cell(row_index, column, cell_for_excel(value))


def _write_summary(
    worksheet: Worksheet,
    session: CleaningSession,
    original: BatchReport,
    current: BatchReport,
) -> None:
    _write_rows(
        worksheet,
        ["Item", "Value"],
        [
            ["Datasets in the working copy", len(session.order)],
            ["Approved operations", len(session.audit)],
            ["Original candidate keys", len(original.candidate_keys)],
            ["Current candidate keys", len(current.candidate_keys)],
            ["Original relationship suggestions", len(original.relationships)],
            ["Current relationship suggestions", len(current.relationships)],
            ["Processing note", PROCESSING_NOTE],
        ],
    )


def _write_inventory(worksheet: Worksheet, session: CleaningSession) -> None:
    rows = []
    for index, dataset_id in enumerate(session.order, start=1):
        item = session.copies[dataset_id]
        rows.append(
            [
                f"Data {index}",
                item.dataset_id,
                item.source_name,
                item.sheet_name or "",
                int(len(item.original)),
                int(len(item.working)),
                int(len(item.working.columns)),
            ]
        )
    _write_rows(
        worksheet,
        ["Sheet", "Dataset", "File", "Worksheet", "Original rows", "Working rows", "Columns"],
        rows,
    )


def _write_metrics(worksheet: Worksheet, original: BatchReport, current: BatchReport) -> None:
    original_items = {item.dataset_id: item for item in original.datasets}
    current_items = {item.dataset_id: item for item in current.datasets}
    rows: list[list[object]] = []
    for dataset_id in list(dict.fromkeys([*original_items, *current_items])):
        before = original_items.get(dataset_id)
        after = current_items.get(dataset_id)
        before_report = before.report if before is not None else None
        after_report = after.report if after is not None else None
        rows.extend(
            [
                [
                    dataset_id,
                    "Data rows",
                    before_report.data_rows if before_report is not None else 0,
                    after_report.data_rows if after_report is not None else 0,
                ],
                [
                    dataset_id,
                    "Rows in exact duplicate groups",
                    before_report.exact_duplicates.rows_in_duplicate_groups.count if before_report else 0,
                    after_report.exact_duplicates.rows_in_duplicate_groups.count if after_report else 0,
                ],
                [
                    dataset_id,
                    "Extra exact copies",
                    before_report.exact_duplicates.extra_copies.count if before_report else 0,
                    after_report.exact_duplicates.extra_copies.count if after_report else 0,
                ],
                [
                    dataset_id,
                    "Missing cells",
                    sum(column.missing.count for column in before_report.columns) if before_report else 0,
                    sum(column.missing.count for column in after_report.columns) if after_report else 0,
                ],
            ]
        )
        names: list[str] = []
        for report in (before_report, after_report):
            if report is None:
                continue
            for column in report.columns:
                if column.name not in names:
                    names.append(column.name)
        before_missing = {column.name: column.missing.count for column in before_report.columns} if before_report else {}
        after_missing = {column.name: column.missing.count for column in after_report.columns} if after_report else {}
        for name in names:
            rows.append(
                [
                    dataset_id,
                    f"Missing values in {name}",
                    before_missing.get(name, 0),
                    after_missing.get(name, 0),
                ]
            )
    _write_rows(worksheet, ["Dataset", "Metric", "Before", "After"], rows)


def _write_relationships(worksheet: Worksheet, original: BatchReport, current: BatchReport) -> None:
    before = {_relationship_key(item): item for item in original.relationships}
    after = {_relationship_key(item): item for item in current.relationships}
    rows: list[list[object]] = []
    for key in sorted(set(before) | set(after)):
        left = before.get(key)
        right = after.get(key)
        sample = right or left
        assert sample is not None
        if left is None:
            status = "suggestion added"
        elif right is None:
            status = "suggestion removed"
        elif left.orphan_count != right.orphan_count:
            status = "orphan count changed"
        else:
            status = "unchanged"
        rows.append(
            [
                sample.parent_dataset,
                sample.parent_column,
                sample.child_dataset,
                sample.child_column,
                "suggestion",
                left.cardinality if left is not None else "",
                right.cardinality if right is not None else "",
                left.orphan_count if left is not None else "",
                right.orphan_count if right is not None else "",
                left.overlap_ratio if left is not None else "",
                right.overlap_ratio if right is not None else "",
                status,
                "Orphan references were not deleted or repaired.",
            ]
        )
    _write_rows(
        worksheet,
        [
            "Parent dataset",
            "Parent column",
            "Child dataset",
            "Child column",
            "Label",
            "Cardinality before",
            "Cardinality after",
            "Orphans before",
            "Orphans after",
            "Overlap before",
            "Overlap after",
            "Status",
            "Note",
        ],
        rows,
    )


def _write_audit(worksheet: Worksheet, audit: list[AuditEntry]) -> None:
    rows = []
    for entry in audit:
        rows.append(
            [
                entry.action_order,
                entry.timestamp,
                entry.dataset_id,
                entry.sheet_name or "",
                entry.operation,
                ", ".join(entry.columns),
                entry.parameter_summary,
                entry.rows_before,
                entry.rows_after,
                entry.rows_removed,
                entry.cells_changed,
                entry.conversion_failures,
                "; ".join(entry.columns_skipped),
                entry.key_warning,
                entry.metrics_before.missing_cells,
                entry.metrics_after.missing_cells,
                entry.metrics_before.extra_duplicate_copies,
                entry.metrics_after.extra_duplicate_copies,
                _relationship_delta(entry.relationships_before, entry.relationships_after),
            ]
        )
    _write_rows(
        worksheet,
        [
            "Order",
            "Time",
            "Dataset",
            "Worksheet",
            "Operation",
            "Columns",
            "Parameters",
            "Rows before",
            "Rows after",
            "Rows removed",
            "Cells changed",
            "Conversion failures",
            "Skipped columns",
            "Key warning",
            "Missing cells before",
            "Missing cells after",
            "Extra copies before",
            "Extra copies after",
            "Relationship changes",
        ],
        rows,
    )


def _relationship_delta(
    before: list[RelationshipSnapshot],
    after: list[RelationshipSnapshot],
) -> str:
    left = {_relationship_key(item): item for item in before}
    right = {_relationship_key(item): item for item in after}
    parts: list[str] = []
    for key in sorted(set(left) | set(right)):
        previous = left.get(key)
        current = right.get(key)
        label = f"{key[0]} · {key[1]} -> {key[2]} · {key[3]}"
        if previous is None and current is not None:
            parts.append(f"{label}: suggestion added, orphans {current.orphan_count}.")
        elif current is None and previous is not None:
            parts.append(f"{label}: suggestion removed, previous orphans {previous.orphan_count}.")
        elif previous is not None and current is not None and previous.orphan_count != current.orphan_count:
            parts.append(f"{label}: orphans {previous.orphan_count} -> {current.orphan_count}.")
    return " ".join(parts) if parts else "No relationship suggestion changed."


def _relationship_key(item: RelationshipSnapshot | object) -> tuple[str, str, str, str]:
    return (
        str(getattr(item, "parent_dataset")),
        str(getattr(item, "parent_column")),
        str(getattr(item, "child_dataset")),
        str(getattr(item, "child_column")),
    )


def _write_rows(worksheet: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    for column, header in enumerate(headers, start=1):
        worksheet.cell(1, column, escape_excel_text(header))
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            if isinstance(value, str):
                stored: object = escape_excel_text(value)
            else:
                stored = value
            worksheet.cell(row_index, column, stored)


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
