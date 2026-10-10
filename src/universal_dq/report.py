"""Export a file profile as JSON, CSV, or Excel without copying cell values."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from universal_dq.models import AnalysisReport, BatchReport, Share


def report_to_json(report: AnalysisReport) -> str:
    """Serialize the profile. The payload has counts and column names only."""
    return report.model_dump_json(indent=2)


def report_to_csv(report: AnalysisReport) -> str:
    """Serialize the same counts as a flat table."""
    buffer = StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=["section", "item", "count", "denominator", "percentage", "status", "detail"],
    )
    writer.writeheader()
    writer.writerow(
        _row(
            "file",
            "data_rows",
            report.data_rows,
            report.data_rows,
            None,
            "",
            report.processing_note,
        )
    )
    writer.writerow(
        _row("file", "column_count", report.column_count, report.column_count, None, "", "")
    )
    exact = report.exact_duplicates
    writer.writerow(_share_row("exact_duplicates", "rows_in_duplicate_groups", exact.rows_in_duplicate_groups, exact.detail))
    writer.writerow(_share_row("exact_duplicates", "extra_copies", exact.extra_copies, exact.detail))
    if report.key_duplicates is not None:
        key = report.key_duplicates
        writer.writerow(
            _share_row(
                "unique_key",
                "rows_in_duplicate_groups",
                key.rows_in_duplicate_groups,
                key.detail,
            )
        )
        writer.writerow(_share_row("unique_key", "extra_copies", key.extra_copies, key.detail))
        writer.writerow(
            _row(
                "unique_key",
                "incomplete_key_rows",
                key.incomplete_key_rows,
                report.data_rows,
                None,
                key.status,
                "Rows excluded from the unique-key denominator because a selected key column is missing.",
            )
        )
    for column in report.columns:
        writer.writerow(
            _share_row(
                "column_missing",
                column.name,
                column.missing,
                f"Inferred type: {column.inferred_type}. Role: {column.missing_role}.",
            )
        )
        if column.mixed_type is not None:
            writer.writerow(
                _row(
                    "mixed_type",
                    column.name,
                    column.mixed_type.non_missing_count,
                    column.mixed_type.non_missing_count,
                    None,
                    column.mixed_type.label,
                    column.mixed_type.detail,
                )
            )
        if column.outlier is not None:
            writer.writerow(
                _row(
                    "outlier",
                    column.name,
                    column.outlier.outlier_count if column.outlier.outlier_count is not None else 0,
                    column.outlier.denominator if column.outlier.denominator is not None else 0,
                    None,
                    column.outlier.label,
                    column.outlier.detail,
                )
            )
    return buffer.getvalue()


def format_share(share: Share) -> str:
    """Render a count with its denominator. Empty denominators stay uncalculated."""
    if share.denominator == 0 or share.ratio is None:
        return f"Not calculated because there are 0 {share.unit}."
    percent = share.ratio * 100
    return f"{share.count:,} of {share.denominator:,} {share.unit} ({percent:.2f}%)."


def _share_row(section: str, item: str, share: Share, detail: str) -> dict[str, object]:
    percentage = "" if share.ratio is None else f"{share.ratio * 100:.2f}%"
    return _row(section, item, share.count, share.denominator, percentage, "", detail)


def _row(
    section: str,
    item: str,
    count: int,
    denominator: int,
    percentage: str | None,
    status: str,
    detail: str,
) -> dict[str, object]:
    return {
        "section": section,
        "item": item,
        "count": count,
        "denominator": denominator,
        "percentage": "" if percentage is None else percentage,
        "status": status,
        "detail": detail,
    }


def report_to_xlsx(
    batch: BatchReport,
    recommendations: list | None = None,
    ai_explanations: list | None = None,
) -> bytes:
    """Build the consolidated workbook in memory. Strings are escaped for Excel.

    Recommendation sheets are added only when ``recommendations`` is provided.
    """
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, batch)
    _write_inventory(workbook.create_sheet("Dataset Inventory"), batch)
    _write_column_profiles(workbook.create_sheet("Column Profiles"), batch)
    _write_missing_values(workbook.create_sheet("Missing Values"), batch)
    _write_duplicate_rows(workbook.create_sheet("Duplicate Rows"), batch)
    _write_key_checks(workbook.create_sheet("Key Checks"), batch)
    _write_mixed_types(workbook.create_sheet("Mixed Types"), batch)
    _write_outliers(workbook.create_sheet("Outliers"), batch)
    _write_candidate_keys(workbook.create_sheet("Candidate Keys"), batch)
    _write_relationships(workbook.create_sheet("Candidate Relationships"), batch)
    _write_limitations(workbook.create_sheet("Limitations"), batch)
    if recommendations is not None:
        from universal_dq.recommendations import append_recommendation_sheets

        append_recommendation_sheets(workbook, recommendations, ai_explanations)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_summary(worksheet: Worksheet, batch: BatchReport) -> None:
    analyzed = [item for item in batch.datasets if item.report is not None]
    failed = [item for item in batch.datasets if item.error is not None]
    data_rows = sum(item.report.data_rows for item in analyzed if item.report is not None)
    _write_table(
        worksheet,
        ["Item", "Value"],
        [
            ["Datasets selected", len(batch.datasets)],
            ["Datasets analyzed", len(analyzed)],
            ["Datasets with errors", len(failed)],
            ["Data rows", data_rows],
            ["Candidate keys", len(batch.candidate_keys)],
            ["Candidate relationships", len(batch.relationships)],
            ["Per-file limit bytes", batch.file_byte_limit],
            ["Total limit bytes", batch.total_byte_limit],
            ["Processing note", batch.processing_note],
        ],
    )


def _write_inventory(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        report = item.report
        rows.append(
            [
                item.dataset_id,
                item.source_name,
                item.sheet_name or "",
                item.file_format or "",
                report.data_rows if report is not None else 0,
                report.column_count if report is not None else 0,
                "analyzed" if report is not None else "error",
                item.error or "",
            ]
        )
    _write_table(
        worksheet,
        ["Dataset", "File", "Worksheet", "Format", "Data rows", "Columns", "Status", "Error"],
        rows,
    )


def _write_column_profiles(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        for column in item.report.columns:
            rows.append(
                [
                    item.dataset_id,
                    item.source_name,
                    item.sheet_name or "",
                    column.name,
                    column.inferred_type,
                    column.missing_role,
                ]
            )
    _write_table(
        worksheet,
        ["Dataset", "File", "Worksheet", "Column", "Inferred type", "Missing label"],
        rows,
    )


def _write_missing_values(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        for column in item.report.columns:
            rows.append(
                [
                    item.dataset_id,
                    item.source_name,
                    item.sheet_name or "",
                    column.name,
                    column.missing.count,
                    column.missing.denominator,
                    _percent(column.missing),
                    column.missing.unit,
                    column.missing_role,
                ]
            )
    _write_table(
        worksheet,
        [
            "Dataset",
            "File",
            "Worksheet",
            "Column",
            "Missing count",
            "Denominator",
            "Percentage",
            "Denominator unit",
            "Label",
        ],
        rows,
    )


def _write_duplicate_rows(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        exact = item.report.exact_duplicates
        rows.append(
            [
                item.dataset_id,
                item.source_name,
                item.sheet_name or "",
                exact.rows_in_duplicate_groups.count,
                exact.rows_in_duplicate_groups.denominator,
                _percent(exact.rows_in_duplicate_groups),
                exact.extra_copies.count,
                exact.extra_copies.denominator,
                _percent(exact.extra_copies),
            ]
        )
    _write_table(
        worksheet,
        [
            "Dataset",
            "File",
            "Worksheet",
            "Rows in duplicate groups",
            "Group denominator",
            "Group percentage",
            "Extra copies",
            "Extra-copy denominator",
            "Extra-copy percentage",
        ],
        rows,
    )


def _write_key_checks(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        key = item.report.key_duplicates
        if key is None:
            rows.append(
                [
                    item.dataset_id,
                    item.source_name,
                    item.sheet_name or "",
                    "",
                    "not selected",
                    0,
                    0,
                    "",
                    "No unique key was selected for this dataset.",
                ]
            )
            continue
        rows.append(
            [
                item.dataset_id,
                item.source_name,
                item.sheet_name or "",
                ", ".join(key.key_columns),
                key.status,
                key.rows_in_duplicate_groups.count,
                key.rows_in_duplicate_groups.denominator,
                _percent(key.rows_in_duplicate_groups),
                key.detail,
            ]
        )
    _write_table(
        worksheet,
        [
            "Dataset",
            "File",
            "Worksheet",
            "Key columns",
            "Status",
            "Rows in duplicate key groups",
            "Denominator",
            "Percentage",
            "Detail",
        ],
        rows,
    )


def _write_mixed_types(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        for column in item.report.columns:
            if column.mixed_type is None:
                continue
            rows.append(
                [
                    item.dataset_id,
                    item.source_name,
                    item.sheet_name or "",
                    column.name,
                    column.mixed_type.label,
                    column.mixed_type.non_missing_count,
                    column.mixed_type.detail,
                ]
            )
    _write_table(
        worksheet,
        ["Dataset", "File", "Worksheet", "Column", "Label", "Non-missing values", "Detail"],
        rows,
    )


def _write_outliers(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = []
    for item in batch.datasets:
        if item.report is None:
            continue
        for column in item.report.columns:
            if column.outlier is None:
                continue
            rows.append(
                [
                    item.dataset_id,
                    item.source_name,
                    item.sheet_name or "",
                    column.name,
                    column.outlier.label,
                    column.outlier.outlier_count if column.outlier.outlier_count is not None else "",
                    column.outlier.denominator if column.outlier.denominator is not None else "",
                    column.outlier.detail,
                ]
            )
    _write_table(
        worksheet,
        ["Dataset", "File", "Worksheet", "Column", "Label", "Outlier count", "Denominator", "Detail"],
        rows,
    )


def _write_candidate_keys(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = [
        [
            item.dataset_id,
            item.column,
            item.data_rows,
            item.distinct_non_missing,
            item.missing_count,
            "suggestion",
            item.detail,
        ]
        for item in batch.candidate_keys
    ]
    _write_table(
        worksheet,
        ["Dataset", "Column", "Data rows", "Distinct values", "Missing values", "Label", "Evidence"],
        rows,
    )


def _write_relationships(worksheet: Worksheet, batch: BatchReport) -> None:
    rows = [
        [
            item.parent_dataset,
            item.parent_column,
            item.child_dataset,
            item.child_column,
            item.cardinality,
            item.matched_count,
            item.child_non_missing,
            item.orphan_count,
            f"{item.overlap_ratio * 100:.2f}%",
            item.label,
            item.detail,
        ]
        for item in batch.relationships
    ]
    _write_table(
        worksheet,
        [
            "Parent dataset",
            "Parent column",
            "Child dataset",
            "Child column",
            "Cardinality",
            "Matched values",
            "Child non-missing",
            "Orphan count",
            "Overlap",
            "Label",
            "Evidence",
        ],
        rows,
    )


def _write_limitations(worksheet: Worksheet, batch: BatchReport) -> None:
    _write_table(worksheet, ["Limitation"], [[item] for item in batch.limitations])


def _write_table(worksheet: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    worksheet.append([_excel_safe(header) for header in headers])
    for row in rows:
        worksheet.append([_excel_safe(value) for value in row])


def _percent(share: Share) -> str:
    if share.ratio is None:
        return "not calculated"
    return f"{share.ratio * 100:.2f}%"


def escape_excel_text(value: object) -> object:
    """Prefix formula-like text. Numbers, dates, and booleans are left unchanged."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _excel_safe(value: object) -> object:
    """Stop Excel from treating untrusted text as a formula."""
    return escape_excel_text(value)
