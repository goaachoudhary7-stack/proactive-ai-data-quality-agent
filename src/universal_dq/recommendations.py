"""Deterministic recommendations from a universal data-quality profile.

Counts come from the profile and the relationship suggestions. This module
does not call a model, does not clean data, and does not treat a mostly
unique column as a business key.
"""

from __future__ import annotations

import hashlib
import json
import re
from typing import Literal, get_args

import pandas as pd
from pydantic import BaseModel, ConfigDict, Field

from universal_dq.models import BatchReport, ColumnProfile
from universal_dq.profile import is_missing
from universal_dq.relationships import _column_tokens

IssueType = Literal[
    "missing_values",
    "exact_duplicates",
    "mixed_types",
    "numerical_outliers",
    "possible_key",
    "orphan_references",
    "formatting",
]
ApprovalKind = Literal["user_approval", "business_rule_confirmation", "both"]
PREFILL_KEY = "dq_clean_prefill"
AI_STORE_KEY = "dq_ai_explanations"

_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_SLASH_DATE = re.compile(r"\d{1,2}/\d{1,2}/\d{4}")
_DASH_DATE = re.compile(r"\d{1,2}-\d{1,2}-\d{4}")
_LETTER = re.compile(r"[A-Za-z]")
_ISSUE_ORDER = {name: index for index, name in enumerate(get_args(IssueType))}


class RemediationOption(BaseModel):
    """A Phase 5C operation the user can open. It is not applied here."""

    model_config = ConfigDict(extra="forbid")

    label: str
    operation: str | None = None
    columns: list[str] = Field(default_factory=list)
    note: str


class Recommendation(BaseModel):
    """One local recommendation. Measurements are separate from hypotheses."""

    model_config = ConfigDict(extra="forbid")

    recommendation_id: str
    source: Literal["deterministic"] = "deterministic"
    issue_type: IssueType
    dataset_id: str
    dataset_ref: str
    column: str | None = None
    column_ref: str | None = None
    sheet_name: str | None = None
    affected_count: int = Field(ge=0)
    denominator: int = Field(ge=0)
    evidence_counts: dict[str, int]
    evidence: str
    why_it_matters: str
    hypotheses: list[str]
    verification_steps: list[str]
    remediation: list[RemediationOption]
    approval: ApprovalKind


class StoredExplanation(BaseModel):
    """An optional explanation. It does not replace confirmed counts."""

    model_config = ConfigDict(extra="forbid")

    recommendation_id: str
    dataset_ref: str
    column_ref: str | None = None
    status: Literal["ready", "error"]
    explanation: str | None = None
    possible_implication: str | None = None
    hypotheses: list[str] = Field(default_factory=list)
    error: str | None = None


def build_recommendations(
    batch: BatchReport,
    frames: dict[str, pd.DataFrame],
) -> list[Recommendation]:
    """Build recommendations for the current profile. Frames are not modified."""
    refs = _refs(batch)
    items: list[Recommendation] = []
    for dataset_id, dataset_ref, columns in refs:
        result = next(item for item in batch.datasets if item.dataset_id == dataset_id)
        report = result.report
        if report is None:
            continue
        frame = frames.get(dataset_id)
        items.extend(_dataset_recommendations(batch, result.sheet_name, dataset_id, dataset_ref, columns, report, frame))
    items.extend(_orphan_recommendations(batch, refs))
    items.sort(key=lambda item: (_ISSUE_ORDER.get(item.issue_type, 99), -item.affected_count, item.recommendation_id))
    return items


def explanation_payload(recommendation: Recommendation) -> dict[str, object]:
    """Summary safe to send after consent. Names and cell values are omitted."""
    payload: dict[str, object] = {
        "dataset_ref": recommendation.dataset_ref,
        "column_ref": recommendation.column_ref,
        "issue_type": recommendation.issue_type,
        "affected_count": recommendation.affected_count,
        "denominator": recommendation.denominator,
        "evidence_counts": dict(recommendation.evidence_counts),
    }
    if recommendation.issue_type == "orphan_references":
        payload["overlap_percent"] = _overlap_percent(recommendation)
    return payload


def evidence_token(recommendations: list[Recommendation]) -> str:
    """Hash the confirmed counts so a changed finding invalidates explanations."""
    body = [
        {
            "id": item.recommendation_id,
            "affected_count": item.affected_count,
            "denominator": item.denominator,
            "evidence_counts": item.evidence_counts,
        }
        for item in recommendations
    ]
    encoded = json.dumps(body, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def explanations_current(
    stored: dict[str, object] | None,
    *,
    fingerprint: str,
    generation: int,
    token: str,
) -> bool:
    """True when stored explanations still match this upload and these counts."""
    if not stored:
        return False
    return (
        stored.get("fingerprint") == fingerprint
        and stored.get("generation") == generation
        and stored.get("token") == token
    )


def _dataset_recommendations(
    batch: BatchReport,
    sheet_name: str | None,
    dataset_id: str,
    dataset_ref: str,
    columns: dict[str, str],
    report: object,
    frame: pd.DataFrame | None,
) -> list[Recommendation]:
    from universal_dq.models import AnalysisReport

    assert isinstance(report, AnalysisReport)
    items: list[Recommendation] = []
    sensitive = _sensitive_columns(batch, dataset_id)
    for column in report.columns:
        column_ref = columns[column.name]
        if column.missing.count > 0:
            items.append(
                _missing(dataset_id, dataset_ref, sheet_name, column, column_ref, report.data_rows, sensitive)
            )
        if column.mixed_type is not None:
            items.append(_mixed(dataset_id, dataset_ref, sheet_name, column, column_ref))
        if (
            column.outlier is not None
            and column.outlier.label == "informational"
            and column.outlier.outlier_count
        ):
            items.append(_outlier(dataset_id, dataset_ref, sheet_name, column, column_ref))
        if frame is not None:
            items.extend(
                _key_gap(dataset_id, dataset_ref, sheet_name, column, column_ref, report.data_rows, frame)
            )
            items.extend(
                _formatting(dataset_id, dataset_ref, sheet_name, column.name, column_ref, frame, sensitive)
            )
    extra = report.exact_duplicates.extra_copies.count
    if extra > 0:
        items.append(_duplicates(dataset_id, dataset_ref, sheet_name, report))
    key = report.key_duplicates
    if key is not None and key.status == "violation":
        items.append(_user_key(dataset_id, dataset_ref, sheet_name, report))
    return items


def _missing(
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column: ColumnProfile,
    column_ref: str,
    data_rows: int,
    sensitive: set[str],
) -> Recommendation:
    options = [
        RemediationOption(
            label="Fill with a constant in the cleaning form",
            operation="fill_constant",
            columns=[column.name],
            note="You type the constant in the cleaning form. Preview and approval are still required.",
        )
    ]
    if column.inferred_type == "number":
        options.append(
            RemediationOption(
                label="Fill numeric gaps with the median in the cleaning form",
                operation="fill_median",
                columns=[column.name],
                note="Median fill runs only when the column is already numeric. Preview and approval are still required.",
            )
        )
    options.append(
        RemediationOption(
            label="Convert whitespace-only text to missing in the cleaning form",
            operation="whitespace_to_missing",
            columns=[column.name],
            note="This stores blank text as missing. Preview and approval are still required.",
        )
    )
    return _item(
        issue_type="missing_values",
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        sheet_name=sheet_name,
        column=column.name,
        column_ref=column_ref,
        affected_count=column.missing.count,
        denominator=data_rows,
        evidence_counts={"missing_count": column.missing.count, "data_rows": data_rows},
        evidence=(
            f"{column.missing.count} of {data_rows} data rows are missing in this column. "
            "Empty cells and whitespace-only text are missing. The text NA and NULL are not."
        ),
        why_it_matters=(
            "Possible implication: missing entries can change totals, averages, and joins. "
            "This is not a measured business loss."
        ),
        hypotheses=[
            "Hypothesis: the value was never collected.",
            "Hypothesis: an upstream export wrote a blank instead of a value.",
            "Hypothesis: a text value made of spaces was loaded as missing.",
        ],
        verification_steps=[
            "Compare the missing count with the column profile on this page.",
            "Check whether blank text and real gaps should be treated the same way.",
        ],
        remediation=options,
        approval=_approval(column.name, sensitive),
    )


def _duplicates(dataset_id: str, dataset_ref: str, sheet_name: str | None, report: object) -> Recommendation:
    from universal_dq.models import AnalysisReport

    assert isinstance(report, AnalysisReport)
    extra = report.exact_duplicates.extra_copies.count
    grouped = report.exact_duplicates.rows_in_duplicate_groups.count
    return _item(
        issue_type="exact_duplicates",
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        sheet_name=sheet_name,
        column=None,
        column_ref=None,
        affected_count=extra,
        denominator=report.data_rows,
        evidence_counts={
            "extra_copies": extra,
            "rows_in_duplicate_groups": grouped,
            "data_rows": report.data_rows,
        },
        evidence=(
            f"{extra} extra exact copies were found beyond the first occurrence. "
            f"{grouped} data rows sit in an exact duplicate group. "
            f"The denominator is {report.data_rows} data rows."
        ),
        why_it_matters=(
            "Possible implication: repeated rows can inflate counts. "
            "This is not a measured business loss."
        ),
        hypotheses=[
            "Hypothesis: the same extract was appended more than once.",
            "Hypothesis: a source system emitted the same record twice.",
        ],
        verification_steps=[
            "Confirm the extra-copy count against the duplicate section of this profile.",
            "Decide whether the first occurrence is the row that should be kept.",
        ],
        remediation=[
            RemediationOption(
                label="Remove exact duplicate rows in the cleaning form",
                operation="drop_exact_duplicates",
                columns=[],
                note="The first matching row is kept. Preview and approval are still required.",
            )
        ],
        approval="user_approval",
    )


def _mixed(
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column: ColumnProfile,
    column_ref: str,
) -> Recommendation:
    assert column.mixed_type is not None
    categories = {kind: int(count) for kind, count in column.mixed_type.categories.items()}
    listed = ", ".join(f"{kind} {count}" for kind, count in sorted(categories.items()))
    return _item(
        issue_type="mixed_types",
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        sheet_name=sheet_name,
        column=column.name,
        column_ref=column_ref,
        affected_count=column.mixed_type.non_missing_count,
        denominator=column.mixed_type.non_missing_count,
        evidence_counts={"non_missing_count": column.mixed_type.non_missing_count, **categories},
        evidence=(
            f"Possible mixed types among {column.mixed_type.non_missing_count} non-missing values: {listed}. "
            "The denominator is the number of non-missing values."
        ),
        why_it_matters=(
            "Possible implication: mixed types can break numeric or date calculations. "
            "This is not a measured business loss."
        ),
        hypotheses=[
            "Hypothesis: some values were stored as text and others as numbers or dates.",
            "Hypothesis: a placeholder word was entered in a numeric column.",
        ],
        verification_steps=[
            "Read the mixed-type category counts in the column profile.",
            "Choose one target type only after checking which category should remain.",
        ],
        remediation=[
            RemediationOption(
                label="Convert the column in the cleaning form",
                operation="convert_type",
                columns=[column.name],
                note="Choose numeric, date, or text in the form. Failed conversions stay unchanged until you preview and approve.",
            )
        ],
        approval="user_approval",
    )


def _outlier(
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column: ColumnProfile,
    column_ref: str,
) -> Recommendation:
    assert column.outlier is not None
    assert column.outlier.outlier_count is not None
    assert column.outlier.denominator is not None
    return _item(
        issue_type="numerical_outliers",
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        sheet_name=sheet_name,
        column=column.name,
        column_ref=column_ref,
        affected_count=column.outlier.outlier_count,
        denominator=column.outlier.denominator,
        evidence_counts={
            "outlier_count": column.outlier.outlier_count,
            "numeric_values": column.outlier.denominator,
        },
        evidence=(
            f"{column.outlier.outlier_count} of {column.outlier.denominator} non-missing numeric values "
            "fall outside the interquartile-range fences. "
            "The denominator is the number of non-missing numeric values."
        ),
        why_it_matters=(
            "Possible implication: extreme values can shift averages. "
            "An extreme value can also be a real measurement. This is not a measured business loss."
        ),
        hypotheses=[
            "Hypothesis: a unit or decimal place differs from the rest of the column.",
            "Hypothesis: the value is a legitimate extreme measurement.",
        ],
        verification_steps=[
            "Compare the outlier count with the column profile. Fence numbers stay on this page.",
            "Decide with the data owner whether any value should be reviewed.",
        ],
        remediation=[
            RemediationOption(
                label="Review only",
                operation=None,
                columns=[column.name],
                note="No cleaning operation removes outliers. Nothing is deleted from this recommendation.",
            )
        ],
        approval="user_approval",
    )


def _user_key(dataset_id: str, dataset_ref: str, sheet_name: str | None, report: object) -> Recommendation:
    from universal_dq.models import AnalysisReport

    assert isinstance(report, AnalysisReport)
    assert report.key_duplicates is not None
    key = report.key_duplicates
    return _item(
        issue_type="possible_key",
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        sheet_name=sheet_name,
        column=None,
        column_ref=None,
        affected_count=key.extra_copies.count,
        denominator=key.extra_copies.denominator,
        evidence_counts={
            "extra_key_copies": key.extra_copies.count,
            "rows_in_duplicate_key_groups": key.rows_in_duplicate_groups.count,
            "incomplete_key_rows": key.incomplete_key_rows,
            "key_column_count": len(key.key_columns),
            "rows_with_complete_key": key.extra_copies.denominator,
        },
        evidence=(
            f"The selected unique-key rule found {key.extra_copies.count} extra copies "
            f"among {key.extra_copies.denominator} rows with a complete key. "
            f"{key.incomplete_key_rows} rows were excluded because the key is incomplete. "
            "This is the user-selected rule, not an inferred business key."
        ),
        why_it_matters=(
            "Possible implication: duplicate keys can attach one record to the wrong parent. "
            "This is not a measured business loss."
        ),
        hypotheses=[
            "Hypothesis: the selected columns are not unique in the source.",
            "Hypothesis: repeated extracts created extra copies of the same key.",
        ],
        verification_steps=[
            "Confirm the selected columns are the intended key.",
            "Compare the extra-copy count with the unique-key section of this profile.",
        ],
        remediation=[
            RemediationOption(
                label="Trim the selected key columns in the cleaning form",
                operation="trim_whitespace",
                columns=list(key.key_columns),
                note="Trim only helps when padding created the difference. Preview and approval are still required.",
            )
        ],
        approval="both",
        subtype="user_rule",
    )


def _key_gap(
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column: ColumnProfile,
    column_ref: str,
    data_rows: int,
    frame: pd.DataFrame,
) -> list[Recommendation]:
    if column.missing.count == 0 or data_rows < 2:
        return []
    tokens = _column_tokens(frame[column.name])
    if len(tokens) < 2 or len(tokens) != len(set(tokens)):
        return []
    return [
        _item(
            issue_type="possible_key",
            dataset_id=dataset_id,
            dataset_ref=dataset_ref,
            sheet_name=sheet_name,
            column=column.name,
            column_ref=column_ref,
            affected_count=column.missing.count,
            denominator=data_rows,
            evidence_counts={
                "missing_count": column.missing.count,
                "distinct_non_missing": len(set(tokens)),
                "data_rows": data_rows,
            },
            evidence=(
                f"{column.missing.count} of {data_rows} data rows are missing. "
                f"The other {len(tokens)} non-missing values are distinct. "
                "Missing values are the only reason this column is not a candidate key. "
                "This is not a confirmed business key."
            ),
            why_it_matters=(
                "Possible implication: an incomplete identifier can drop rows from a join. "
                "This is not a measured business loss."
            ),
            hypotheses=[
                "Hypothesis: the missing cells are unknown identifiers rather than duplicates.",
                "Hypothesis: the column is not used as an identifier at all.",
            ],
            verification_steps=[
                "Confirm with the data owner whether this column is intended to identify a row.",
                "Do not treat the column as a key until that rule is confirmed.",
            ],
            remediation=[
                RemediationOption(
                    label="Fill missing values in the cleaning form only after the rule is confirmed",
                    operation="fill_constant",
                    columns=[column.name],
                    note="A filled value does not prove the column is a key. Preview and approval are still required.",
                )
            ],
            approval="both",
            subtype="missing_blocks_candidate",
        )
    ]


def _formatting(
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column_name: str,
    column_ref: str,
    frame: pd.DataFrame,
    sensitive: set[str],
) -> list[Recommendation]:
    padded = 0
    shapes: dict[str, int] = {}
    cases: dict[str, int] = {"upper": 0, "lower": 0, "mixed": 0}
    letter_values = 0
    for value in frame[column_name].tolist():
        if is_missing(value) or not isinstance(value, str):
            continue
        stripped = value.strip()
        if stripped != value and stripped != "":
            padded += 1
        shape = _date_shape(stripped)
        if shape is not None:
            shapes[shape] = shapes.get(shape, 0) + 1
        if _LETTER.search(stripped):
            letter_values += 1
            cases[_case_style(stripped)] += 1
    items: list[Recommendation] = []
    if padded > 0:
        items.append(
            _item(
                issue_type="formatting",
                dataset_id=dataset_id,
                dataset_ref=dataset_ref,
                sheet_name=sheet_name,
                column=column_name,
                column_ref=column_ref,
                affected_count=padded,
                denominator=int(len(frame)),
                evidence_counts={"padded_text_count": padded, "data_rows": int(len(frame))},
                evidence=(
                    f"{padded} text values have leading or trailing spaces and still contain other characters. "
                    f"The denominator is {len(frame)} data rows."
                ),
                why_it_matters=(
                    "Possible implication: padded text can fail an exact match. "
                    "This is not a measured business loss."
                ),
                hypotheses=[
                    "Hypothesis: the export padded a fixed-width field.",
                    "Hypothesis: a value was pasted with a surrounding space.",
                ],
                verification_steps=[
                    "Check the padded-text count before trimming.",
                    "Trim only the columns you select in the cleaning form.",
                ],
                remediation=[
                    RemediationOption(
                        label="Trim whitespace in the cleaning form",
                        operation="trim_whitespace",
                        columns=[column_name],
                        note="Preview and approval are still required.",
                    )
                ],
                approval=_approval(column_name, sensitive),
                subtype="whitespace",
            )
        )
    used_cases = {name: count for name, count in cases.items() if count > 0}
    if letter_values >= 4 and len(used_cases) >= 2:
        items.append(
            _item(
                issue_type="formatting",
                dataset_id=dataset_id,
                dataset_ref=dataset_ref,
                sheet_name=sheet_name,
                column=column_name,
                column_ref=column_ref,
                affected_count=letter_values,
                denominator=int(len(frame)),
                evidence_counts={"letter_values": letter_values, **used_cases, "data_rows": int(len(frame))},
                evidence=(
                    f"{letter_values} text values contain letters. "
                    + ", ".join(f"{name} {count}" for name, count in sorted(used_cases.items()))
                    + f". The denominator is {len(frame)} data rows."
                ),
                why_it_matters=(
                    "Possible implication: mixed letter case can split a group that was meant to be one label. "
                    "This is not a measured business loss."
                ),
                hypotheses=[
                    "Hypothesis: different sources used different capitalization.",
                    "Hypothesis: the mixed case is intentional, such as a code and a name.",
                ],
                verification_steps=[
                    "Confirm whether one letter case is required.",
                    "Choose upper, lower, or title case in the cleaning form.",
                ],
                remediation=[
                    RemediationOption(
                        label="Standardize letter case in the cleaning form",
                        operation="change_case",
                        columns=[column_name],
                        note="You choose the case in the form. Preview and approval are still required.",
                    )
                ],
                approval=_approval(column_name, sensitive),
                subtype="letter_case",
            )
        )
    if len(shapes) >= 2:
        items.append(
            _item(
                issue_type="formatting",
                dataset_id=dataset_id,
                dataset_ref=dataset_ref,
                sheet_name=sheet_name,
                column=column_name,
                column_ref=column_ref,
                affected_count=sum(shapes.values()),
                denominator=int(len(frame)),
                evidence_counts={**shapes, "data_rows": int(len(frame))},
                evidence=(
                    "More than one date shape is present: "
                    + ", ".join(f"{name} {count}" for name, count in sorted(shapes.items()))
                    + f". The denominator is {len(frame)} data rows."
                ),
                why_it_matters=(
                    "Possible implication: mixed date shapes can be read as different days. "
                    "This is not a measured business loss."
                ),
                hypotheses=[
                    "Hypothesis: one extract used year-month-day text and another used slashes.",
                    "Hypothesis: the shapes are codes rather than dates.",
                ],
                verification_steps=[
                    "Choose an explicit date format. Ambiguous dates are not guessed.",
                    "Values that do not match the format stay unchanged.",
                ],
                remediation=[
                    RemediationOption(
                        label="Convert to a date in the cleaning form",
                        operation="convert_type",
                        columns=[column_name],
                        note="Enter the date format in the form. Preview and approval are still required.",
                    )
                ],
                approval=_approval(column_name, sensitive),
                subtype="date_shape",
            )
        )
    return items


def _orphan_recommendations(
    batch: BatchReport,
    refs: list[tuple[str, str, dict[str, str]]],
) -> list[Recommendation]:
    lookup = {dataset_id: (dataset_ref, columns) for dataset_id, dataset_ref, columns in refs}
    sheets = {item.dataset_id: item.sheet_name for item in batch.datasets}
    items: list[Recommendation] = []
    for link in batch.relationships:
        if link.orphan_count <= 0:
            continue
        parent = lookup.get(link.parent_dataset)
        child = lookup.get(link.child_dataset)
        if parent is None or child is None:
            continue
        parent_ref, parent_columns = parent
        child_ref, child_columns = child
        overlap_percent = int(round(link.overlap_ratio * 100))
        items.append(
            _item(
                issue_type="orphan_references",
                dataset_id=link.child_dataset,
                dataset_ref=child_ref,
                sheet_name=sheets.get(link.child_dataset),
                column=link.child_column,
                column_ref=child_columns.get(link.child_column),
                affected_count=link.orphan_count,
                denominator=link.child_non_missing,
                evidence_counts={
                    "orphan_count": link.orphan_count,
                    "matched_count": link.matched_count,
                    "child_non_missing": link.child_non_missing,
                    "overlap_percent": overlap_percent,
                },
                evidence=(
                    f"{link.orphan_count} non-missing child values do not occur in the suggested parent column. "
                    f"{link.matched_count} of {link.child_non_missing} non-missing child values match. "
                    "This is a suggestion, not a confirmed foreign key. The unmatched values are not listed."
                ),
                why_it_matters=(
                    "Possible implication: unmatched references can be dropped from a join. "
                    "This is not a measured business loss."
                ),
                hypotheses=[
                    "Hypothesis: the child value was never created in the parent dataset.",
                    "Hypothesis: padding or letter case kept equal values from matching.",
                ],
                verification_steps=[
                    "Review the relationship suggestion and the orphan count on this page.",
                    "Do not delete unmatched rows from this recommendation.",
                ],
                remediation=[
                    RemediationOption(
                        label="Trim the child column in the cleaning form",
                        operation="trim_whitespace",
                        columns=[link.child_column],
                        note="Trim only helps when spaces caused the mismatch. Preview and approval are still required. Orphan rows are not deleted.",
                    )
                ],
                approval="both",
                subtype=f"{parent_ref}:{parent_columns.get(link.parent_column, '')}:{child_ref}",
            )
        )
    return items


def _item(
    *,
    issue_type: IssueType,
    dataset_id: str,
    dataset_ref: str,
    sheet_name: str | None,
    column: str | None,
    column_ref: str | None,
    affected_count: int,
    denominator: int,
    evidence_counts: dict[str, int],
    evidence: str,
    why_it_matters: str,
    hypotheses: list[str],
    verification_steps: list[str],
    remediation: list[RemediationOption],
    approval: ApprovalKind,
    subtype: str = "",
) -> Recommendation:
    identity = "|".join([issue_type, dataset_ref, column_ref or "-", subtype])
    return Recommendation(
        recommendation_id=identity,
        issue_type=issue_type,
        dataset_id=dataset_id,
        dataset_ref=dataset_ref,
        column=column,
        column_ref=column_ref,
        sheet_name=sheet_name,
        affected_count=affected_count,
        denominator=denominator,
        evidence_counts=evidence_counts,
        evidence=evidence,
        why_it_matters=why_it_matters,
        hypotheses=hypotheses,
        verification_steps=verification_steps,
        remediation=remediation,
        approval=approval,
    )


def _refs(batch: BatchReport) -> list[tuple[str, str, dict[str, str]]]:
    refs: list[tuple[str, str, dict[str, str]]] = []
    index = 1
    for item in batch.datasets:
        if item.report is None:
            continue
        columns = {name: f"column_{position}" for position, name in enumerate(item.report.column_names, start=1)}
        refs.append((item.dataset_id, f"dataset_{index}", columns))
        index += 1
    return refs


def _sensitive_columns(batch: BatchReport, dataset_id: str) -> set[str]:
    names = {item.column for item in batch.candidate_keys if item.dataset_id == dataset_id}
    for link in batch.relationships:
        if link.parent_dataset == dataset_id:
            names.add(link.parent_column)
        if link.child_dataset == dataset_id:
            names.add(link.child_column)
    return names


def _approval(column: str, sensitive: set[str]) -> ApprovalKind:
    if column in sensitive:
        return "both"
    return "user_approval"


def _date_shape(value: str) -> str | None:
    if _ISO_DATE.fullmatch(value):
        return "year_month_day"
    if _SLASH_DATE.fullmatch(value):
        return "slash_date"
    if _DASH_DATE.fullmatch(value):
        return "dash_date"
    return None


def _case_style(value: str) -> str:
    letters = "".join(character for character in value if character.isalpha())
    if letters.isupper():
        return "upper"
    if letters.islower():
        return "lower"
    return "mixed"


def _overlap_percent(recommendation: Recommendation) -> int:
    return int(recommendation.evidence_counts.get("overlap_percent", 0))


def append_recommendation_sheets(
    workbook: object,
    recommendations: list[Recommendation],
    explanations: list[StoredExplanation] | None = None,
) -> None:
    """Add the two recommendation sheets. The caller has already built the profile sheets."""
    from openpyxl.workbook.workbook import Workbook

    from universal_dq.report import escape_excel_text

    assert isinstance(workbook, Workbook)
    _write_recommendation_sheet(workbook, recommendations, escape_excel_text)
    _write_explanation_sheet(workbook, recommendations, explanations or [], escape_excel_text)


def _write_recommendation_sheet(workbook: object, recommendations: list[Recommendation], escape) -> None:
    worksheet = workbook.create_sheet("Recommendations")  # type: ignore[attr-defined]
    headers = [
        "Source",
        "Issue",
        "Dataset",
        "Dataset ref",
        "Worksheet",
        "Column",
        "Column ref",
        "Affected count",
        "Denominator",
        "Confirmed evidence",
        "Possible implication",
        "Hypotheses",
        "Verification",
        "Suggested action",
        "Approval",
    ]
    rows: list[list[object]] = []
    if not recommendations:
        rows.append(["deterministic", "none", "", "", "", "", "", 0, 0, "No recommendation met the evidence bar.", "", "", "", "", ""])
    for item in recommendations:
        actions = " ".join(option.label for option in item.remediation)
        rows.append(
            [
                item.source,
                item.issue_type,
                item.dataset_id,
                item.dataset_ref,
                item.sheet_name or "",
                item.column or "",
                item.column_ref or "",
                item.affected_count,
                item.denominator,
                item.evidence,
                item.why_it_matters,
                " ".join(item.hypotheses),
                " ".join(item.verification_steps),
                actions,
                item.approval,
            ]
        )
    _write(worksheet, headers, rows, escape)


def _write_explanation_sheet(
    workbook: object,
    recommendations: list[Recommendation],
    explanations: list[StoredExplanation],
    escape,
) -> None:
    worksheet = workbook.create_sheet("AI Explanations")  # type: ignore[attr-defined]
    headers = [
        "Status",
        "Dataset ref",
        "Column ref",
        "Confirmed affected count",
        "AI explanation",
        "Possible implication from AI",
        "AI hypotheses",
        "Note",
    ]
    by_id = {item.recommendation_id: item for item in recommendations}
    rows: list[list[object]] = []
    if not explanations:
        rows.append(
            [
                "not_requested",
                "",
                "",
                "",
                "",
                "",
                "",
                "No AI explanation was requested. Confirmed counts stay on the Recommendations sheet.",
            ]
        )
    for item in explanations:
        confirmed = by_id.get(item.recommendation_id)
        rows.append(
            [
                item.status,
                item.dataset_ref,
                item.column_ref or "",
                confirmed.affected_count if confirmed is not None else "",
                item.explanation or item.error or "",
                item.possible_implication or "",
                " ".join(item.hypotheses),
                "This text is an AI explanation. It does not replace the confirmed counts.",
            ]
        )
    _write(worksheet, headers, rows, escape)


def _write(worksheet: object, headers: list[str], rows: list[list[object]], escape) -> None:
    for column, header in enumerate(headers, start=1):
        worksheet.cell(1, column, escape(header))  # type: ignore[attr-defined]
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            stored = escape(value) if isinstance(value, str) else value
            worksheet.cell(row_index, column, stored)  # type: ignore[attr-defined]
