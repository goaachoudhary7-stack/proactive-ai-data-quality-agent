"""Tests for approved cleaning of uploaded tables.

These tests build frames in memory. They do not call a model, start Streamlit,
or read the banking ground-truth manifest.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from universal_dq.batch import LoadedDataset, build_batch_report, collect_datasets  # noqa: E402
from universal_dq.cleaning import (  # noqa: E402
    CleaningRejected,
    apply_operation,
    content_fingerprint,
    datasets_for_quality,
    new_cleaning_session,
    preview_operation,
    reset,
    restore_or_create,
    undo,
)
from universal_dq.cleaning_export import (  # noqa: E402
    AUDIT_SHEETS,
    cleaned_dataset_csv,
    cleaned_datasets_xlsx,
    cleaning_audit_xlsx,
)
from universal_dq.models import OperationSpec  # noqa: E402
from universal_dq.report import report_to_xlsx  # noqa: E402

WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
SECRET_FILL = "SECRET-FILL-VALUE"
SECRET_CELL = "RAW-SECRET-CELL"
ORPHAN = "ORPHAN-TOKEN"


def _frame(rows: dict[str, list[object]]) -> pd.DataFrame:
    return pd.DataFrame(rows, dtype=object)


def _session(frames: dict[str, pd.DataFrame]):
    datasets = [
        LoadedDataset(dataset_id, dataset_id, "csv", None, frame, None)
        for dataset_id, frame in frames.items()
    ]
    return new_cleaning_session(datasets)


def _spec(dataset_id: str, operation: str, **kwargs) -> OperationSpec:
    return OperationSpec(operation=operation, dataset_id=dataset_id, **kwargs)  # type: ignore[arg-type]


def _apply(session, spec: OperationSpec):
    return apply_operation(session, spec, approved=True, now=WHEN)


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        worksheet = workbook.active if first else workbook.create_sheet(name)
        worksheet.title = name
        first = False
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_strings(content: bytes) -> list[str]:
    workbook = load_workbook(BytesIO(content))
    found: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    found.append(cell.value)
    return found


def test_each_cleaning_operation_counts_rows_and_cells() -> None:
    session = _session({"people.csv": _frame({"Code": ["  ab ", "cd"], "City": ["  x ", "y"]})})
    original_id = id(session.copies["people.csv"].original)
    trimmed = _apply(session, _spec("people.csv", "trim_whitespace", columns=["Code"]))
    assert trimmed.copies["people.csv"].working["Code"].tolist() == ["ab", "cd"]
    assert trimmed.copies["people.csv"].working["City"].tolist() == ["  x ", "y"]
    assert trimmed.audit[0].cells_changed == 1
    assert id(trimmed.copies["people.csv"].original) == original_id

    blanks = _session({"people.csv": _frame({"Code": ["NA", "NULL", "  ", ""]})})
    blanked = _apply(blanks, _spec("people.csv", "whitespace_to_missing", columns=["Code"]))
    assert blanked.copies["people.csv"].working["Code"].tolist() == ["NA", "NULL", None, None]
    assert blanked.audit[0].cells_changed == 2

    duplicates = _session({"people.csv": _frame({"Code": ["A", "A", "B"], "City": ["x", "x", "y"]})})
    deduped = _apply(duplicates, _spec("people.csv", "drop_exact_duplicates"))
    assert deduped.audit[0].rows_removed == 1
    assert deduped.copies["people.csv"].working["Code"].tolist() == ["A", "B"]

    emptied = _session({"people.csv": _frame({"A": [1, None, 0], "B": ["x", "  ", False]})})
    removed = _apply(emptied, _spec("people.csv", "drop_empty_rows"))
    assert removed.audit[0].rows_removed == 1
    assert removed.copies["people.csv"].working["A"].tolist() == [1, 0]
    assert removed.copies["people.csv"].working["B"].tolist() == ["x", False]

    missing = _session({"people.csv": _frame({"Code": ["A", None]})})
    filled = _apply(missing, _spec("people.csv", "fill_constant", columns=["Code"], fill_value="Unknown"))
    assert filled.copies["people.csv"].working["Code"].tolist() == ["A", "Unknown"]
    assert filled.audit[0].cells_changed == 1
    assert "Unknown" not in filled.audit[0].parameter_summary

    amounts = _session({"people.csv": _frame({"Amount": [1, 2, 3, None]})})
    median = _apply(amounts, _spec("people.csv", "fill_median", columns=["Amount"]))
    assert median.copies["people.csv"].working["Amount"].tolist() == [1, 2, 3, 2]
    assert median.audit[0].parameter_summary.endswith("The median is not stored in this audit.")

    titled = _apply(session, _spec("people.csv", "change_case", columns=["Code"], case="title"))
    assert titled.copies["people.csv"].working["Code"].tolist()[0] == "  Ab "

    quantities = _session({"people.csv": _frame({"Qty": ["8", "no"]})})
    numeric = _apply(quantities, _spec("people.csv", "convert_type", columns=["Qty"], target_type="numeric"))
    assert numeric.copies["people.csv"].working["Qty"].tolist() == [8, "no"]
    assert numeric.audit[0].conversion_failures == 1
    assert numeric.audit[0].cells_changed == 1

    dates = _session({"people.csv": _frame({"When": ["02/03/2024", "2024-01-09"]})})
    dated = _apply(
        dates,
        _spec("people.csv", "convert_type", columns=["When"], target_type="date", date_format="%m/%d/%Y"),
    )
    assert dated.copies["people.csv"].working["When"].tolist() == [datetime(2024, 2, 3), "2024-01-09"]
    assert dated.audit[0].conversion_failures == 1

    renamed = _apply(session, _spec("people.csv", "rename_columns", renames={"City": "Town"}))
    assert list(renamed.copies["people.csv"].working.columns) == ["Code", "Town"]
    pd.testing.assert_frame_equal(session.copies["people.csv"].working, session.copies["people.csv"].original)


def test_preview_does_not_change_frames_and_apply_requires_approval() -> None:
    session = _session({"people.csv": _frame({"Code": ["  ab ", "cd"]})})
    spec = _spec("people.csv", "trim_whitespace", columns=["Code"])
    before = session.copies["people.csv"].working.copy()
    preview = preview_operation(session, spec)
    assert preview.applicable is True
    assert preview.cells_changed == 1
    pd.testing.assert_frame_equal(session.copies["people.csv"].working, before)
    with pytest.raises(CleaningRejected):
        apply_operation(session, spec, approved=False, now=WHEN)
    pd.testing.assert_frame_equal(session.copies["people.csv"].working, before)


def test_dates_require_an_explicit_format_and_keep_failures() -> None:
    session = _session({"dates.csv": _frame({"When": ["02/03/2024", "02/30/2024"]})})
    missing_format = preview_operation(
        session,
        _spec("dates.csv", "convert_type", columns=["When"], target_type="date"),
    )
    assert missing_format.applicable is False
    assert missing_format.block_reason is not None
    assert "guessed" in missing_format.block_reason

    iso = _apply(
        session,
        _spec("dates.csv", "convert_type", columns=["When"], target_type="date", date_format="%Y-%m-%d"),
    )
    assert iso.copies["dates.csv"].working["When"].tolist() == ["02/03/2024", "02/30/2024"]
    assert iso.audit[0].conversion_failures == 2

    month_first = _apply(
        session,
        _spec("dates.csv", "convert_type", columns=["When"], target_type="date", date_format="%m/%d/%Y"),
    )
    assert month_first.copies["dates.csv"].working["When"].tolist()[0] == datetime(2024, 2, 3)
    assert month_first.copies["dates.csv"].working["When"].tolist()[1] == "02/30/2024"
    assert month_first.audit[0].conversion_failures == 1

    day_first = _apply(
        session,
        _spec("dates.csv", "convert_type", columns=["When"], target_type="date", date_format="%d/%m/%Y"),
    )
    assert day_first.copies["dates.csv"].working["When"].tolist()[0] == datetime(2024, 3, 2)
    assert session.copies["dates.csv"].original["When"].tolist() == ["02/03/2024", "02/30/2024"]


def test_numeric_conversion_does_not_treat_booleans_as_numbers() -> None:
    session = _session(
        {"flags.csv": _frame({"Flag": [True, False, True], "Amount": ["10", "ten", None]})}
    )
    flags = _apply(session, _spec("flags.csv", "convert_type", columns=["Flag"], target_type="numeric"))
    assert flags.audit[0].conversion_failures == 3
    assert flags.copies["flags.csv"].working["Flag"].tolist() == [True, False, True]
    amounts = _apply(session, _spec("flags.csv", "convert_type", columns=["Amount"], target_type="numeric"))
    assert amounts.copies["flags.csv"].working["Amount"].tolist() == [10, "ten", None]
    assert amounts.audit[0].conversion_failures == 1
    text = _apply(session, _spec("flags.csv", "convert_type", columns=["Flag"], target_type="text"))
    assert text.copies["flags.csv"].working["Flag"].tolist() == ["True", "False", "True"]


def test_median_fill_skips_columns_that_are_not_already_numeric() -> None:
    session = _session(
        {
            "numbers.csv": _frame(
                {
                    "Amount": [1, 2, None],
                    "Label": ["1", "2", None],
                    "Flag": [True, False, None],
                }
            )
        }
    )
    mixed = preview_operation(session, _spec("numbers.csv", "fill_median", columns=["Label"]))
    assert mixed.applicable is False
    assert mixed.block_reason is not None
    assert "numeric" in mixed.block_reason
    flags = preview_operation(session, _spec("numbers.csv", "fill_median", columns=["Flag"]))
    assert flags.applicable is False
    applied = _apply(session, _spec("numbers.csv", "fill_median", columns=["Amount", "Label"]))
    assert applied.copies["numbers.csv"].working["Amount"].tolist() == [1, 2, 1.5]
    assert applied.copies["numbers.csv"].working["Label"].tolist() == ["1", "2", None]
    assert applied.audit[0].columns_skipped


def test_rename_rejects_blank_and_duplicate_names() -> None:
    session = _session({"people.csv": _frame({"A": [1], "B": [2]})})
    duplicate = preview_operation(session, _spec("people.csv", "rename_columns", renames={"A": "B"}))
    blank = preview_operation(session, _spec("people.csv", "rename_columns", renames={"A": "  "}))
    assert duplicate.applicable is False
    assert blank.applicable is False
    assert list(session.copies["people.csv"].working.columns) == ["A", "B"]


def test_undo_and_reset_restore_working_copies_without_touching_originals() -> None:
    session = _session({"people.csv": _frame({"City": ["  aa ", "bb"]})})
    original = session.copies["people.csv"].original
    trimmed = _apply(session, _spec("people.csv", "trim_whitespace", columns=["City"]))
    upper = _apply(trimmed, _spec("people.csv", "change_case", columns=["City"], case="upper"))
    assert upper.copies["people.csv"].working["City"].tolist() == ["AA", "BB"]
    once = undo(upper)
    assert once.copies["people.csv"].working["City"].tolist() == ["aa", "bb"]
    assert once.audit[0].operation == "trim_whitespace"
    assert once.audit[0].timestamp == trimmed.audit[0].timestamp
    twice = undo(once)
    pd.testing.assert_frame_equal(twice.copies["people.csv"].working, original)
    assert twice.audit == []
    assert undo(twice) is twice
    restored = reset(upper)
    pd.testing.assert_frame_equal(restored.copies["people.csv"].working, original)
    assert restored.copies["people.csv"].working is not original
    assert restored.copies["people.csv"].original is original
    assert restored.audit == []
    pd.testing.assert_frame_equal(session.copies["people.csv"].working, original)


def test_datasets_and_worksheets_are_cleaned_independently() -> None:
    left = _frame({"City": ["  a ", "b"]})
    right = _frame({"City": ["  c ", "d"]})
    session = _session({"team.csv": left, "team.csv (2)": right})
    right_working = session.copies["team.csv (2)"].working
    updated = _apply(session, _spec("team.csv", "trim_whitespace", columns=["City"]))
    assert updated.copies["team.csv"].working["City"].tolist() == ["a", "b"]
    assert updated.copies["team.csv (2)"].working is right_working
    pd.testing.assert_frame_equal(updated.copies["team.csv (2)"].working, right)

    content = _xlsx(
        {
            "People": [["City"], ["  a "], ["b"]],
            "Orders": [["City"], ["  c "], ["d"]],
        }
    )
    loaded = collect_datasets([("staff.xlsx", content, ("People", "Orders"))])
    workbook_session = new_cleaning_session(loaded)
    assert workbook_session.order == ["staff.xlsx · People", "staff.xlsx · Orders"]
    cleaned = _apply(
        workbook_session,
        _spec("staff.xlsx · People", "trim_whitespace", columns=["City"]),
    )
    assert cleaned.copies["staff.xlsx · People"].working["City"].tolist() == ["a", "b"]
    assert cleaned.copies["staff.xlsx · Orders"].working["City"].tolist() == ["  c ", "d"]

    repeated = collect_datasets(
        [
            ("team.csv", b"City\n  a \nb\n", None),
            ("team.csv", b"City\n  c \nd\n", None),
        ]
    )
    assert [item.dataset_id for item in repeated] == ["team.csv", "team.csv (2)"]


def test_content_fingerprint_uses_file_bytes() -> None:
    first = b"A,B\n1,2\n"
    second = b"A,B\n9,8\n"
    assert len(first) == len(second)
    first_fp = content_fingerprint([("same.csv", first, None)])
    second_fp = content_fingerprint([("same.csv", second, None)])
    assert first_fp != second_fp
    assert content_fingerprint([("same.csv", first, None)]) == first_fp
    assert content_fingerprint([("same.csv", first, ("People",))]) != content_fingerprint(
        [("same.csv", first, ("Orders",))]
    )
    loaded_first = collect_datasets([("same.csv", first, None)])
    loaded_second = collect_datasets([("same.csv", second, None)])
    session = new_cleaning_session(loaded_first)
    assert restore_or_create(session, first_fp, first_fp, loaded_second) is session
    replaced = restore_or_create(session, first_fp, second_fp, loaded_second)
    assert replaced is not session
    assert replaced.copies["same.csv"].original.iloc[0, 0] == "9"


def test_audit_omits_fill_constants_and_cell_values() -> None:
    session = _session({"people.csv": _frame({"Amount": [SECRET_CELL, None]})})
    updated = _apply(
        session,
        _spec("people.csv", "fill_constant", columns=["Amount"], fill_value=SECRET_FILL),
    )
    payload = json.dumps([entry.model_dump() for entry in updated.audit])
    assert SECRET_FILL not in payload
    assert SECRET_CELL not in payload
    assert updated.copies["people.csv"].working["Amount"].tolist() == [SECRET_CELL, SECRET_FILL]
    original = build_batch_report(datasets_for_quality(updated, which="original"))
    current = build_batch_report(datasets_for_quality(updated, which="working"))
    audit_book = cleaning_audit_xlsx(updated, original, current)
    audit_text = "\n".join(_cell_strings(audit_book))
    assert SECRET_FILL not in audit_text
    assert SECRET_CELL not in audit_text
    assert "Cleaning Audit" in load_workbook(BytesIO(audit_book)).sheetnames
    cleaned_text = "\n".join(_cell_strings(cleaned_datasets_xlsx(updated)))
    assert SECRET_FILL in cleaned_text
    assert SECRET_CELL in cleaned_text
    undone = undo(updated)
    assert undone.copies["people.csv"].working["Amount"].tolist() == [SECRET_CELL, None]


def test_excel_export_preserves_typed_values_and_escapes_text() -> None:
    frame = _frame(
        {
            "=Total": [42, 7],
            "Ratio": [1.5, 2.5],
            "Flag": [True, False],
            "When": [datetime(2024, 5, 6, 7, 8, 9), datetime(2024, 1, 2)],
            "Note": ["=1+1", "ok"],
        }
    )
    session = _session({"people.csv": frame})
    workbook = load_workbook(BytesIO(cleaned_datasets_xlsx(session)))
    sheet = workbook["Data 1"]
    assert sheet["A1"].value == "'=Total"
    assert sheet["A2"].value == 42
    assert sheet["A2"].data_type == "n"
    assert sheet["B2"].value == 1.5
    assert sheet["B2"].data_type == "n"
    assert sheet["C2"].value is True
    assert sheet["C2"].data_type == "b"
    assert sheet["D2"].value == datetime(2024, 5, 6, 7, 8, 9)
    assert sheet["D2"].data_type == "d"
    assert sheet["E2"].value == "'=1+1"
    csv_text = cleaned_dataset_csv(frame).decode("utf-8-sig")
    assert "'=Total" in csv_text
    assert "'=1+1" in csv_text
    assert "42" in csv_text
    assert "'42" not in csv_text

    formula_session = _session({"=evil.csv": _frame({"=Code": ["  SECRET-CELL"]})})
    cleaned = _apply(formula_session, _spec("=evil.csv", "trim_whitespace", columns=["=Code"]))
    audit_strings = _cell_strings(
        cleaning_audit_xlsx(
            cleaned,
            build_batch_report(datasets_for_quality(cleaned, which="original")),
            build_batch_report(datasets_for_quality(cleaned, which="working")),
        )
    )
    assert "SECRET-CELL" not in "\n".join(audit_strings)
    assert "'=Code" in audit_strings
    assert "'=evil.csv" in audit_strings
    assert all(not value.startswith(("=", "+", "-", "@")) for value in audit_strings)
    assert load_workbook(BytesIO(cleaning_audit_xlsx(
        cleaned,
        build_batch_report(datasets_for_quality(cleaned, which="original")),
        build_batch_report(datasets_for_quality(cleaned, which="working")),
    ))).sheetnames == list(AUDIT_SHEETS)


def test_relationship_orphan_counts_are_recalculated_and_rows_stay() -> None:
    parent = _frame({"Employee Code": ["E100", "E200", "E300", "E400"]})
    child = _frame(
        {
            "Employee Code": ["e100", "e100", "e200", "e300", ORPHAN.lower()],
            "Line": ["1", "2", "3", "4", "5"],
        }
    )
    session = _session({"people.csv": parent, "orders.csv": child})
    before = build_batch_report(datasets_for_quality(session, which="original"))
    assert before.relationships == []
    spaced = _session({"people.csv": _frame({"Employee Code": [" E100", " E200", " E300", " E400"]})})
    warning = preview_operation(
        spaced,
        _spec("people.csv", "trim_whitespace", columns=["Employee Code"]),
    )
    assert warning.applicable is True
    assert any("candidate primary key" in note for note in warning.warnings)
    assert any("will not be deleted" in note for note in warning.warnings)

    updated = _apply(session, _spec("orders.csv", "change_case", columns=["Employee Code"], case="upper"))
    assert updated.copies["orders.csv"].working["Employee Code"].tolist()[-1] == ORPHAN
    assert len(updated.copies["orders.csv"].working) == 5
    assert updated.copies["people.csv"].working is session.copies["people.csv"].working
    assert updated.audit[0].relationships_before == []
    assert updated.audit[0].relationships_after[0].orphan_count == 1
    assert updated.audit[0].relationships_after[0].cardinality == "one-to-many"
    payload = updated.audit[0].model_dump_json()
    assert ORPHAN not in payload
    assert ORPHAN.lower() not in payload
    current = build_batch_report(datasets_for_quality(updated, which="working"))
    assert current.relationships[0].orphan_count == 1
    original_book = report_to_xlsx(before)
    assert "Candidate Relationships" in load_workbook(BytesIO(original_book)).sheetnames
    assert updated.audit[0].metrics_before.extra_duplicate_copies == 0


def test_duplicate_removal_updates_before_and_after_metrics() -> None:
    session = _session({"people.csv": _frame({"Code": ["A", "A", "A"], "City": ["x", "x", "x"]})})
    preview = preview_operation(session, _spec("people.csv", "drop_exact_duplicates"))
    assert preview.rows_removed == 2
    updated = _apply(session, _spec("people.csv", "drop_exact_duplicates"))
    assert updated.audit[0].metrics_before.extra_duplicate_copies == 2
    assert updated.audit[0].metrics_after.extra_duplicate_copies == 0
    assert updated.audit[0].metrics_before.rows_in_duplicate_groups == 3
    assert updated.copies["people.csv"].working["Code"].tolist() == ["A"]
    original = build_batch_report(datasets_for_quality(updated, which="original"))
    assert original.datasets[0].report is not None
    assert original.datasets[0].report.exact_duplicates.extra_copies.count == 2


def test_invalid_requests_leave_the_working_copy_unchanged() -> None:
    session = _session({"people.csv": _frame({"Code": []})})
    empty = preview_operation(session, _spec("people.csv", "drop_exact_duplicates"))
    missing = preview_operation(session, _spec("people.csv", "trim_whitespace", columns=["Missing"]))
    blank_fill = preview_operation(
        session,
        _spec("people.csv", "fill_constant", columns=["Code"], fill_value=""),
    )
    assert empty.applicable is False
    assert missing.applicable is False
    assert blank_fill.applicable is False
    assert session.copies["people.csv"].working.empty


def test_numeric_text_and_numbers_follow_the_existing_duplicate_rule() -> None:
    numbers = _session({"people.csv": _frame({"A": [1, 1.0, "1"]})})
    removed = _apply(numbers, _spec("people.csv", "drop_exact_duplicates"))
    assert removed.audit[0].rows_removed == 1
    assert removed.copies["people.csv"].working["A"].tolist() == [1, "1"]
    text = _session({"people.csv": _frame({"A": ["1", "1.0"]})})
    kept = preview_operation(text, _spec("people.csv", "drop_exact_duplicates"))
    assert kept.applicable is False
