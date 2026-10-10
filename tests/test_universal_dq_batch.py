"""Tests for multi-file analysis, suggested relationships, and the Excel report."""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from universal_dq.batch import build_batch_report, collect_datasets  # noqa: E402
from universal_dq.report import report_to_xlsx  # noqa: E402

REPORT_SHEETS = {
    "Summary",
    "Dataset Inventory",
    "Column Profiles",
    "Missing Values",
    "Duplicate Rows",
    "Key Checks",
    "Mixed Types",
    "Outliers",
    "Candidate Keys",
    "Candidate Relationships",
    "Limitations",
}


def _csv(text: str) -> bytes:
    return text.encode("utf-8")


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    names = list(sheets)
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for name, rows in sheets.items():
        worksheet = workbook[name]
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _batch(files: list[tuple[str, bytes, tuple[str, ...] | None]], **kwargs):
    datasets = collect_datasets(files, **kwargs)
    return build_batch_report(datasets), datasets


def _cell_strings(content: bytes) -> list[str]:
    workbook = load_workbook(BytesIO(content))
    found: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    found.append(value)
    return found


def test_multiple_files_and_duplicate_filenames_are_analyzed_separately() -> None:
    first = _csv("Employee Code,Department\nE1,HR\nE2,HR\n")
    second = _csv("Employee Code,Department\nE1,Ops\nE3,Ops\n")
    batch, _datasets = _batch([("team.csv", first, None), ("team.csv", second, None)])
    assert [item.dataset_id for item in batch.datasets] == ["team.csv", "team.csv (2)"]
    assert all(item.report is not None and item.report.data_rows == 2 for item in batch.datasets)
    assert {item.column for item in batch.candidate_keys} == {"Employee Code"}


def test_selected_worksheets_become_separate_datasets() -> None:
    content = _xlsx(
        {
            "People": [
                ["Employee Code", "Department", "Note"],
                ["E100", "HR", "SECRET-SHEET-CELL"],
                ["E200", "Sales", "kept"],
                ["E300", "HR", "kept"],
                ["E400", "Ops", "kept"],
            ],
            "Orders": [
                ["Employee Code", "Product"],
                ["E100", "Pen"],
                ["E100", "Book"],
                ["E200", "Desk"],
                ["E300", "Lamp"],
                ["ORPHAN-TOKEN-77", "Chair"],
            ],
        }
    )
    batch, datasets = _batch([("staff.xlsx", content, ("People", "Orders"))])
    assert [item.dataset_id for item in datasets] == [
        "staff.xlsx · People",
        "staff.xlsx · Orders",
    ]
    assert len(batch.relationships) == 1
    relationship = batch.relationships[0]
    assert relationship.label == "suggestion"
    assert relationship.cardinality == "one-to-many"
    assert relationship.parent_dataset == "staff.xlsx · People"
    assert relationship.parent_column == "Employee Code"
    assert relationship.child_column == "Employee Code"
    assert relationship.matched_count == 4
    assert relationship.child_non_missing == 5
    assert relationship.orphan_count == 1
    assert relationship.overlap_ratio == 0.8
    payload = json.dumps(batch.model_dump(mode="json"))
    workbook = report_to_xlsx(batch)
    assert "SECRET-SHEET-CELL" not in payload
    assert "ORPHAN-TOKEN-77" not in payload
    assert "SECRET-SHEET-CELL" not in _cell_strings(workbook)
    assert "ORPHAN-TOKEN-77" not in _cell_strings(workbook)
    assert set(load_workbook(BytesIO(workbook)).sheetnames) == REPORT_SHEETS


def test_unreadable_file_does_not_drop_the_other_file() -> None:
    good = _csv("Badge Number,City\nB1,East\nB2,West\n")
    batch, _datasets = _batch(
        [
            ("broken.xlsx", b"this is not a workbook", None),
            ("offices.csv", good, None),
        ]
    )
    assert batch.datasets[0].error is not None
    assert "could not be opened" in batch.datasets[0].error
    assert batch.datasets[1].report is not None
    assert batch.datasets[1].report.data_rows == 2


def test_total_and_per_file_limits_skip_only_the_oversized_file() -> None:
    small = b"A,B\n1,2\n3,4\n"
    batch, datasets = _batch(
        [("one.csv", small, None), ("two.csv", small, None)],
        max_total_bytes=len(small),
    )
    assert datasets[0].error is None
    assert datasets[1].error is not None
    assert "total" in datasets[1].error
    assert batch.datasets[0].report is not None
    limited, _ignored = _batch([("wide.csv", small, None)], max_file_bytes=4)
    assert limited.datasets[0].error is not None
    assert "per-file" in limited.datasets[0].error


def test_relationships_are_not_inferred_without_enough_evidence() -> None:
    numbers = _csv("Badge Number\n1\n2\n3\n4\n")
    quantities = _csv("Quantity\n1\n2\n3\n4\n")
    different, _datasets = _batch(
        [("badges.csv", numbers, None), ("counts.csv", quantities, None)]
    )
    assert different.relationships == []

    parent = _csv("Employee Code\nE1\nE2\nE3\nE4\n")
    child = _csv("Employee Code\nE1\nZ2\nZ3\nZ4\nZ5\n")
    weak, _datasets = _batch([("parent.csv", parent, None), ("child.csv", child, None)])
    assert weak.relationships == []
    assert any(item.column == "Employee Code" for item in weak.candidate_keys)


def test_excel_report_escapes_formula_like_text() -> None:
    content = _csv("=2+2,City\nA,East\nB,West\n")
    batch, _datasets = _batch([("=evil.csv", content, None)])
    values = _cell_strings(report_to_xlsx(batch))
    assert any(value == "'=2+2" for value in values)
    assert any(value == "'=evil.csv" for value in values)
    assert not any(value.startswith(("=", "+", "-", "@")) for value in values)
