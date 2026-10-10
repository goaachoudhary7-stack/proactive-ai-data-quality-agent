"""Tests for deterministic recommendations and optional explanations.

These tests build tables in memory. They do not call a model or read the
banking ground-truth manifest.
"""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from ai_agents.upload_explainer import (  # noqa: E402
    claims_confirmed_cause,
    explain_recommendation,
    gemini_narrative_source,
    unsupported_numbers,
    validate_narrative,
)
from universal_dq.batch import LoadedDataset, build_batch_report  # noqa: E402
from universal_dq.cleaning import apply_operation, new_cleaning_session  # noqa: E402
from universal_dq.models import OperationSpec  # noqa: E402
from universal_dq.recommendations import (  # noqa: E402
    StoredExplanation,
    build_recommendations,
    evidence_token,
    explanation_payload,
    explanations_current,
)
from universal_dq.report import report_to_xlsx  # noqa: E402

SECRET = "SECRET-CELL"
FILENAME = "secret-file.csv"


def _frame(rows: dict[str, list[object]]) -> pd.DataFrame:
    return pd.DataFrame(rows, dtype=object)


def _batch(frames: dict[str, pd.DataFrame], **kwargs):
    loaded = [
        LoadedDataset(dataset_id, dataset_id, "csv", None, frame, None)
        for dataset_id, frame in frames.items()
    ]
    batch = build_batch_report(loaded, **kwargs)
    return batch, {item.dataset_id: item.frame for item in loaded if item.frame is not None}


def _ids(recommendations, issue_type: str) -> list[str]:
    return [item.column or item.dataset_id for item in recommendations if item.issue_type == issue_type]


def test_recommendations_use_confirmed_counts_for_each_issue() -> None:
    frame = _frame(
        {
            "Code": ["E1", "E2", None, "E3", "E4", "E5", "E6", "E7"],
            "Amount": [1, 2, 3, 4, 5, 6, 7, 100],
            "Label": [1, "x", 2, 3, 4, 5, 6, 7],
            "City": ["  alpha ", "ALPHA", "beta", "BETA", "gamma", "GAMMA", "delta", "DELTA"],
            "When": [
                "2024-01-02",
                "02/03/2024",
                "02-03-2024",
                "04/05/2024",
                "2024-06-07",
                "08/09/2024",
                "10-11-2024",
                "12/01/2024",
            ],
        }
    )
    duplicated = _frame({"Code": ["A", "A"], "City": ["x", "x"]})
    batch, frames = _batch({"people.csv": frame, "copies.csv": duplicated})
    before = frames["people.csv"].copy()
    recommendations = build_recommendations(batch, frames)
    pd.testing.assert_frame_equal(frames["people.csv"], before)

    missing = next(item for item in recommendations if item.issue_type == "missing_values" and item.column == "Code")
    assert missing.affected_count == 1
    assert missing.denominator == 8
    assert missing.evidence_counts["missing_count"] == 1
    assert all(option.operation != "drop_exact_duplicates" or option.columns == [] for option in missing.remediation)
    assert any(option.operation == "fill_constant" for option in missing.remediation)

    duplicates = next(item for item in recommendations if item.issue_type == "exact_duplicates")
    assert duplicates.affected_count == 1
    assert duplicates.remediation[0].operation == "drop_exact_duplicates"

    mixed = next(item for item in recommendations if item.issue_type == "mixed_types")
    assert mixed.column == "Label"
    assert mixed.evidence_counts["text"] == 1
    assert mixed.remediation[0].operation == "convert_type"

    outlier = next(item for item in recommendations if item.issue_type == "numerical_outliers")
    assert outlier.column == "Amount"
    assert outlier.affected_count >= 1
    assert outlier.remediation[0].operation is None

    key = next(
        item
        for item in recommendations
        if item.issue_type == "possible_key" and item.column == "Code"
    )
    assert "not a confirmed business key" in key.evidence
    assert key.approval == "both"

    formatting = [item.recommendation_id for item in recommendations if item.issue_type == "formatting"]
    assert any(item.endswith("whitespace") for item in formatting)
    assert any(item.endswith("letter_case") for item in formatting)
    assert any(item.endswith("date_shape") for item in formatting)
    assert SECRET not in json.dumps([item.model_dump() for item in recommendations])


def test_high_uniqueness_is_not_treated_as_a_business_key() -> None:
    values = ["A", "A"] + [f"K{index}" for index in range(18)]
    batch, frames = _batch({"people.csv": _frame({"Code": values})})
    recommendations = build_recommendations(batch, frames)
    assert [item.issue_type for item in recommendations if item.issue_type == "possible_key"] == []


def test_user_selected_key_violation_is_recommended() -> None:
    batch, frames = _batch(
        {"people.csv": _frame({"Code": ["A", "A", "B"], "City": ["x", "y", "z"]})},
        key_columns=["Code"],
    )
    recommendations = build_recommendations(batch, frames)
    key = next(item for item in recommendations if item.issue_type == "possible_key")
    assert key.evidence_counts["extra_key_copies"] == 1
    assert "user-selected rule" in key.evidence
    assert key.remediation[0].operation == "trim_whitespace"
    assert key.remediation[0].columns == ["Code"]


def test_orphan_recommendation_requires_an_existing_suggestion() -> None:
    parent = _frame({"Employee Code": ["E100", "E200", "E300", "E400"]})
    child = _frame(
        {
            "Employee Code": ["E100", "E100", "E200", "E300", "ORPHAN-TOKEN"],
            "Line": ["1", "2", "3", "4", "5"],
        }
    )
    weak = _frame({"Quantity": [1, 2, 3, 4]})
    batch, frames = _batch({"people.csv": parent, "orders.csv": child, "other.csv": weak})
    recommendations = build_recommendations(batch, frames)
    orphans = [item for item in recommendations if item.issue_type == "orphan_references"]
    assert len(orphans) == 1
    assert orphans[0].affected_count == 1
    assert orphans[0].evidence_counts["matched_count"] == 4
    assert "ORPHAN-TOKEN" not in orphans[0].evidence
    assert orphans[0].remediation[0].operation == "trim_whitespace"
    assert all(item.dataset_id != "other.csv" or item.issue_type != "orphan_references" for item in recommendations)


def test_payload_omits_filenames_cell_values_and_column_names() -> None:
    frame = _frame({"Notes": [SECRET, None]})
    batch, frames = _batch({FILENAME: frame})
    recommendations = build_recommendations(batch, frames)
    payload = explanation_payload(recommendations[0])
    encoded = json.dumps(payload)
    assert SECRET not in encoded
    assert FILENAME not in encoded
    assert "Notes" not in encoded
    assert payload["dataset_ref"] == "dataset_1"
    assert payload["column_ref"] == "column_1"
    assert payload["affected_count"] == 1
    assert "fill" not in encoded


def test_cleaning_recalculates_recommendations_without_applying_them() -> None:
    frame = _frame({"City": ["  aa ", "bb", "cc", "dd"]})
    batch, frames = _batch({"people.csv": frame})
    before = build_recommendations(batch, frames)
    assert any(item.issue_type == "formatting" and item.recommendation_id.endswith("whitespace") for item in before)
    session = new_cleaning_session(
        [LoadedDataset("people.csv", "people.csv", "csv", None, frame, None)]
    )
    updated = apply_operation(
        session,
        OperationSpec(operation="trim_whitespace", dataset_id="people.csv", columns=["City"]),
        approved=True,
    )
    current, current_frames = _batch({"people.csv": updated.copies["people.csv"].working})
    after = build_recommendations(current, current_frames)
    assert not any(item.recommendation_id.endswith("whitespace") for item in after)
    assert updated.copies["people.csv"].original["City"].tolist()[0] == "  aa "


def test_consent_invalid_output_and_service_failures() -> None:
    payload = {
        "dataset_ref": "dataset_1",
        "column_ref": "column_1",
        "issue_type": "missing_values",
        "affected_count": 4,
        "denominator": 10,
        "evidence_counts": {"missing_count": 4, "data_rows": 10},
    }

    def fail_if_called(_payload):
        raise AssertionError("source called")

    blocked = explain_recommendation(payload, consent=False, narrative_source=fail_if_called)
    assert blocked.status == "error"
    assert blocked.explanation is None

    accepted = validate_narrative(
        payload,
        {
            "explanation": "4 of 10 rows are missing in dataset_1.",
            "possible_implication": "Possible effect on a total.",
            "hypotheses": ["Hypothesis: the value was not collected."],
            "cause_status": "hypothesis",
        },
    )
    assert accepted.explanation.startswith("4 of 10")
    assert unsupported_numbers("4 of 10", payload) == []
    assert not claims_confirmed_cause("This is not a confirmed root cause.")

    try:
        validate_narrative(
            payload,
            {
                "explanation": "99999 rows are missing.",
                "possible_implication": "Possible effect.",
                "hypotheses": ["Hypothesis: a gap."],
                "cause_status": "hypothesis",
            },
        )
        raise AssertionError("invented count was accepted")
    except ValueError as exc:
        assert "number" in str(exc)

    try:
        validate_narrative(
            payload,
            {
                "explanation": "This is a confirmed root cause.",
                "possible_implication": "Possible effect.",
                "hypotheses": ["Hypothesis: a gap."],
                "cause_status": "hypothesis",
            },
        )
        raise AssertionError("confirmed cause was accepted")
    except ValueError as exc:
        assert "confirmed" in str(exc)

    def raise_timeout(_payload):
        raise TimeoutError("timed out")

    def raise_quota(_payload):
        raise RuntimeError("429 quota exceeded")

    def raise_key(_payload):
        raise RuntimeError("API key is not configured.")

    held = [dict(payload)]
    for source in (raise_timeout, raise_quota, raise_key):
        result = explain_recommendation(payload, consent=True, narrative_source=source)
        assert result.status == "error"
        assert result.explanation is None
        assert held == [payload]


def test_explanations_become_stale_when_the_findings_change() -> None:
    stored = {"fingerprint": "aaa", "generation": 1, "token": "same", "items": {"one": "text"}}
    assert explanations_current(stored, fingerprint="aaa", generation=1, token="same")
    assert not explanations_current(stored, fingerprint="bbb", generation=1, token="same")
    assert not explanations_current(stored, fingerprint="aaa", generation=2, token="same")
    assert not explanations_current(stored, fingerprint="aaa", generation=1, token="other")
    batch, frames = _batch({"people.csv": _frame({"City": ["  aa ", "bb", "cc", "dd"]})})
    first = build_recommendations(batch, frames)
    assert evidence_token(first) != evidence_token([])


def test_recommendation_workbook_is_separate_from_the_original_report() -> None:
    batch, frames = _batch({FILENAME: _frame({"Notes": [SECRET, None], "=Code": ["=1+1", "ok"]})})
    original = report_to_xlsx(batch)
    original_names = set(load_workbook(BytesIO(original)).sheetnames)
    assert "Recommendations" not in original_names
    assert "AI Explanations" not in original_names
    recommendations = build_recommendations(batch, frames)
    explanation = StoredExplanation(
        recommendation_id=recommendations[0].recommendation_id,
        dataset_ref="dataset_1",
        column_ref="column_1",
        status="ready",
        explanation="MODEL-NOTE-77",
        possible_implication="Possible implication only.",
        hypotheses=["Hypothesis: a gap."],
    )
    workbook = load_workbook(BytesIO(report_to_xlsx(batch, recommendations, [explanation])))
    assert original_names <= set(workbook.sheetnames)
    assert "Recommendations" in workbook.sheetnames
    assert "AI Explanations" in workbook.sheetnames
    texts = []
    for name in workbook.sheetnames:
        for row in workbook[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    texts.append((name, cell.value))
    joined = "\n".join(value for _name, value in texts)
    assert SECRET not in joined
    assert "MODEL-NOTE-77" in joined
    assert all(name == "AI Explanations" for name, value in texts if "MODEL-NOTE-77" in value)
    assert "'=Code" in joined or any(value.startswith("'=") for _name, value in texts)
    assert all(not value.startswith(("=", "+", "-", "@")) for _name, value in texts)


class _FakeModels:
    def __init__(self, response=None, error=None) -> None:
        self.response = response
        self.error = error
        self.calls: list[dict[str, object]] = []

    def generate_content(self, *, model, contents, config):
        self.calls.append({"model": model, "contents": contents, "config": config})
        if self.error is not None:
            raise self.error
        return self.response


class _FakeClient:
    def __init__(self, response=None, error=None) -> None:
        self.models = _FakeModels(response, error)


class _ApiError(Exception):
    def __init__(self, code: int, message: str) -> None:
        super().__init__(message)
        self.code = code


def _summary() -> dict[str, object]:
    return {
        "dataset_ref": "dataset_1",
        "column_ref": "column_1",
        "issue_type": "missing_values",
        "affected_count": 4,
        "denominator": 10,
        "evidence_counts": {"missing_count": 4, "data_rows": 10},
    }


def _ready_response():
    return type(
        "Response",
        (),
        {
            "parsed": {
                "explanation": "4 of 10 rows are missing in dataset_1.",
                "possible_implication": "Possible effect on a total.",
                "hypotheses": ["Hypothesis: the value was not collected."],
                "cause_status": "hypothesis",
            },
            "text": None,
        },
    )()


def _explain_with(client: _FakeClient, payload: dict[str, object] | None = None):
    body = payload or _summary()

    def source(received):
        return gemini_narrative_source(
            received,
            client=client,
            model="gemini/gemini-3.8-flash",
        )

    return explain_recommendation(body, consent=True, narrative_source=source), client


def test_direct_gemini_request_sends_the_approved_summary_once() -> None:
    payload = _summary()
    result, client = _explain_with(_FakeClient(_ready_response()), payload)
    assert result.status == "ready"
    assert result.explanation == "4 of 10 rows are missing in dataset_1."
    assert result.hypotheses == ["Hypothesis: the value was not collected."]
    assert len(client.models.calls) == 1
    call = client.models.calls[0]
    contents = str(call["contents"])
    assert call["model"] == "gemini-3.8-flash"
    assert json.dumps(payload, indent=2, sort_keys=True) in contents
    assert SECRET not in contents
    assert FILENAME not in contents
    assert "Notes" not in contents
    config = call["config"]
    assert isinstance(config, dict)
    assert config["response_mime_type"] == "application/json"
    assert config["response_schema"].__name__ == "AiNarrative"
    assert config["automatic_function_calling"] == {"disable": True}
    assert config["thinking_config"] == {"include_thoughts": False, "thinking_level": "MINIMAL"}
    assert payload == _summary()


def test_consent_blocks_the_direct_gemini_request() -> None:
    client = _FakeClient(_ready_response())

    def source(received):
        return gemini_narrative_source(received, client=client, model="gemini/gemini-3.8-flash")

    blocked = explain_recommendation(_summary(), consent=False, narrative_source=source)
    assert blocked.status == "error"
    assert blocked.explanation is None
    assert client.models.calls == []


def test_direct_gemini_errors_keep_the_offline_summary() -> None:
    payload = _summary()
    cases = (
        (_ApiError(401, "UNAUTHENTICATED"), "API key"),
        (_ApiError(404, "NOT_FOUND"), "unavailable"),
        (_ApiError(429, "RESOURCE_EXHAUSTED"), "quota"),
        (TimeoutError("timed out"), "timed out"),
        (type("ReadTimeout", (Exception,), {})("The read operation timed out"), "timed out"),
    )
    for error, expected in cases:
        result, client = _explain_with(_FakeClient(error=error), dict(payload))
        assert result.status == "error"
        assert result.explanation is None
        assert expected.casefold() in (result.error or "").casefold()
        assert len(client.models.calls) == 1
        assert payload == _summary()


def test_malformed_gemini_text_is_discarded() -> None:
    response = type("Response", (), {"parsed": None, "text": "not json"})()
    result, _client = _explain_with(_FakeClient(response))
    assert result.status == "error"
    assert result.explanation is None
    assert "required explanation shape" in (result.error or "")


def test_missing_key_is_reported_without_building_a_request(monkeypatch) -> None:
    class Settings:
        model = "gemini/gemini-3.8-flash"
        api_key_env = "DQ_TEST_EXPLANATION_KEY"

    monkeypatch.setattr("ai_agents.investigator.resolve_llm_settings", lambda: Settings())
    monkeypatch.delenv("DQ_TEST_EXPLANATION_KEY", raising=False)
    try:
        gemini_narrative_source(_summary())
        raise AssertionError("missing key was accepted")
    except RuntimeError as exc:
        assert "API key is not configured" in str(exc)


def test_non_gemini_model_is_not_sent() -> None:
    client = _FakeClient(_ready_response())
    try:
        gemini_narrative_source(_summary(), client=client, model="openai/gpt-4o")
        raise AssertionError("non-gemini model was accepted")
    except RuntimeError as exc:
        assert "not found" in str(exc)
    assert client.models.calls == []


def test_explainer_does_not_use_a_crew() -> None:
    source = (SRC / "ai_agents" / "upload_explainer.py").read_text(encoding="utf-8")
    assert "crewai" not in source.casefold()
    assert "kickoff" not in source
    assert ".csv" not in source
